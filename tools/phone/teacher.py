from pathlib import Path
import os
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
R=Path(__file__).resolve().parents[2];W=Path(os.environ.get('HWT_BUILD_WORKSPACE',R.parent));O=R/'du/textbook/2hcl/unit3/shou-ji-bei-diu-qi-zhi-hou';L=json.loads((O/'lesson.json').read_text());B=json.loads((O/'teaching-brief.json').read_text());T=L['tasks']
D=Document();sec=D.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=D.styles[name];st.font.name='Noto Serif CJK SC';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt({'Normal':11,'Title':24,'Subtitle':12,'Heading 1':16,'Heading 2':13,'Heading 3':12}[name]);st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=Pt(18 if name=='Normal' else 24)
D.styles['Title'].paragraph_format.space_after=Pt(14)
def p(s='',bold=False):
 a=D.add_paragraph();a.add_run(s).bold=bold;return a
def lead(a,b):
 q=D.add_paragraph();q.add_run(a+'：').bold=True;q.add_run(b);return q
def h(s,level=1):D.add_heading(s,level)
def page():D.add_page_break()
def table(headers,rows,widths):
 tb=D.add_table(rows=1,cols=len(headers));tb.autofit=False
 for i,w in enumerate(widths):tb.columns[i].width=Inches(w)
 for i,a in enumerate(headers):tb.rows[0].cells[i].text=str(a)
 for row in rows:
  for i,a in enumerate(row):tb.add_row() if False else None
  cs=tb.add_row().cells
  for i,a in enumerate(row):cs[i].text=str(a)
 for ri,row in enumerate(tb.rows):
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
   for edge in ['top','bottom','left','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');bd.append(x)
   pr.append(bd);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE7EB'if ri==0 else'FFFFFF');pr.append(sh);mar=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:w'),'90');x.set(qn('w:type'),'dxa');mar.append(x)
   pr.append(mar)
   for q in c.paragraphs:
    q.paragraph_format.line_spacing=Pt(16);q.paragraph_format.space_after=Pt(2)
    for r in q.runs:r.font.size=Pt(10.5);r.bold=ri==0
 tb.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'));return tb

def item(i,full=True):
 t=T[i];h(t['title'],2)
 for q in t['questions']:p(q,True)
 lead('来源与课堂身份',t['source']+'；'+t['role'])
 lead('参考答案',t['answer']);lead('证据与加工',t['explain']);lead('分层提示',' → '.join(t['hints']));lead('常见失分',t['reject'])
 if full:lead('评分建议',t['scoreGuide'])
 for q,a in t.get('oral',[]):lead(q,a)

D.add_paragraph(L['title'],style='Title');D.add_paragraph('中二高级华文 单元三 阅读巩固 教师流程与参考答案',style='Subtitle')
p('本课用两课时引导学生结合图表读懂电子垃圾的构成、增长、回收价值与危害，并用证据解释说明方法的作用。按9个自然段推进，10组任务分成11个小步骤；所有已提供课本问题和教参引导题均已对应。')
h('核心目标')
for i,g in enumerate(B['goals'],1):p(str(i)+'　'+g)
h('两课时安排')
for r in B['schedule']:lead(r[0]+' '+r[1],r[2])
lead('课件分课','共36页。第一课第1—17页，第二课第18—36页；问题先出现，答案后揭。')
lead('书写负担','第一课只写数字作用及一句仿句；第二课只写说明方法及开放回应。其他题以点选、配对、口头解释为主。')
lead('我做 我们做 你做','教师示范第1段作用题；师生共读第2段构成和第6段对照；学生独立完成数字作用、仿句、说明方法与开放回应。')
page();h('段落 图画与问题对应')
table(['段落','说明重点','原图格号'],[[i+1,L['paragraphIdeas'][i],'、'.join(map(str,L['panels'][i]))]for i in range(9)],[.55,4.95,1.5])
h('主要理解障碍',2)
for s in B['difficulties']:p(s)
lead('阅读顺序','左边完整原文，右边本段图画及词语积累；题目页左文右题。仿句随第3段处理；读完第9段后才做全文开放回应。')
page();h('教材数字与原图边界')
for s in B['sourceBoundary']:p(s)
h('教师核对原则',2)
p('本文件的参考答案和评分建议依据教材、教参整理，不冒充官方考试评分细则。等义答案按意义单位接受；可独立成立的理解分别给分。引用原句本身不扣分，作用题还须解释原句支持什么重点。')
p('未增加新短文、课后一周小测、课时、打卡或竞赛。教参中的课后作业页码及外部资料搜索建议不转成新增必做任务。')
page();h('第一课 第一段与第二段');item(0,False);item(1,False)
page();h('第一课 第三段图表与仿句');item(2)
lead('教材表格对应','2014年4440万吨；2019年比2014年增加21%；预计2030年约7400万吨。说明重点是电子垃圾总量庞大并持续增长，问题日益严重。')
lead('数字核对','六类构成32＋24＋20＋13.5＋9＋1.5＝100%；家用电器合计32＋24＋20＝76%。2030相对2014为7400÷4440≈1.67倍，增加约67%。不要求学生推算教材未标出的逐年精确值。')
page();h('第一课 仿句与第四段原因');item(3);item(4,False)
page();h('第二课 第五段说明方法');item(5)
lead('完整取证示范','第5段先加入“一部智能手机的材料通常包含60种金属元素，里头的贵重金属，如黄金、白银等可以再循环利用。”再加入“数据显示，如果把2016年废弃手机中的金属全部回收，价值可达95亿美元。”两句保留后，学生写自己的分析。')
page();h('第二课 第六段与第七段');item(6,False);item(7,False)
h('意义单位与边界',2)
p('第6段需要把“价值高”和“实际利用少”联系起来，不只报数字。第7段按“填埋→重金属渗透→污染土壤和地下水→健康受影响”组织；回答可以简洁，但不能用“影响环境”四字替代具体危害。')
page();h('第二课 第八段与第九段');item(8,False);item(9,False)
h('直接信息与推断',2)
p('第8段原文明说集中收集、专人处理、再循环利用；“方便公众”是依据全岛广设回收点作出的合理推断。第9段联系前文的原因、资源与危害来解释劝告，不能只重复“三思而行”四个字。')
page();h('必做开放回应');item(10)
for a in T[10]['alternatives']:p(a)
h('建议评分的独立性',2)
p('具体建议、合理理由、吻合例子、解释例子与环保效果的联系，各1分。只写口号不得当成充分理由；不同于示例的建议，只要能说明因果且合理可行，均可接受。不用字数或关键词自动判分。')
page();h('课堂操作与记录')
for label,txt in [('学生入口','reading.html。读文、第一课、第二课、检查、个人词语本；课本原题及来源可见。'),('答题反馈','第一次答错先给定位提示，5秒后可订正；再错后可看讲解。主观题先填写后展示参考，保留首次回答与订正。'),('多句证据','点左边原句，再点“加入原句”；重复点同句不重复添加，后一句不会覆盖前一句。可粘贴、编辑及换行。'),('个人词语本','点击原文已收录词或自行输入词语均可收藏。未收录词建卡但标释义待补充；可编辑、导入导出及翻卡复习。'),('保存范围','姓名、班级、学号／组别、作答、首次记录、尝试和订正保存在当前设备和浏览器。检查页导出JSON、CSV和文字记录；没有账号云同步或教师自动收件。'),('跨课文收藏','沿用同域共享词语本，按年级、单元、课文分类。保留已有笔记、释义、出处和熟悉状态；不强制加入全部预设词条。'),('离线后备','教师页下载独立HTML或完整ZIP。ZIP解压后打开“开始使用.html”。平台词语预习需联网；朗读依赖设备语音包。'),('资料下载','PPT文字和数据图表可编辑；原有图画及教材折线图为图像。课件字体使用霞鹜文楷与思源宋体；未安装字体的设备可能替换。在线投影使用课件渲染图片。'),('公开范围','学生入口与教师资料分开，但教师页为公开页面，不提供身份保护。'),('验证范围','完成教学核对、课件与文档渲染、资源及DOM交互检查。未做实体iPad及原生PowerPoint实机测试。')]:lead(label,txt)
h('自检结果',2)
p('已核对题干与证据、中文优先、段落次序、多句取证、必做开放回应、选择题四个正确位置各一次、图表比例合计与课时总长。第3段近似翻倍的表述已保留并另作精确计算说明。')
# page fields
foot=sec.footer.paragraphs[0];foot.alignment=2;r=foot.add_run('手机被丢弃之后  ');r.font.size=Pt(9);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
for root in [D.element,D.styles.element]:
 for border in root.xpath('.//w:pBdr'):border.getparent().remove(border)
D.save(O/'downloads/teacher-guide.docx')
print('Teacher guide ready')
