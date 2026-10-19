from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2]; D=ROOT/'du/textbook/2hcl/unit1/fang-feng-zheng';HERE=Path(__file__).parent
L=json.loads((HERE/'lesson.json').read_text()); M=json.loads((HERE/'teaching_meta.json').read_text());S=json.loads((D/'slides.json').read_text())
doc=Document(); sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for b in doc.styles.element.xpath('.//w:pBdr'): b.getparent().remove(b)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[name];s.font.name='Noto Serif CJK SC';s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');s.font.color.rgb=RGBColor.from_string('243D37')
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=Pt(16);doc.styles['Normal'].paragraph_format.space_after=Pt(5)
for name,size in [('Title',25),('Heading 1',18),('Heading 2',13),('Heading 3',12)]:doc.styles[name].font.size=Pt(size)
for name in ['Heading 1','Heading 2','Heading 3']:doc.styles[name].paragraph_format.space_before=Pt(9);doc.styles[name].paragraph_format.space_after=Pt(5)
footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('华文通 · 放风筝 · 教师版　');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for r in footer.runs:r.font.size=Pt(9)
def p(s,bold=False,kai=False):
    a=doc.add_paragraph();r=a.add_run(str(s));r.bold=bold
    if kai:r.font.name='LXGW WenKai TC';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'LXGW WenKai TC')
    return a
def h(s,level=1):doc.add_heading(s,level)
def new(s):doc.add_page_break();h(s)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,v in enumerate(headers):t.rows[0].cells[i].text=str(v)
    for vals in rows:
        c=t.add_row().cells
        for i,v in enumerate(vals):c[i].text=str(v)
    head=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(head)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);pr=c._tc.get_or_add_tcPr()
            if ri==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E7EFE9');pr.append(fill)
            borders=OxmlElement('w:tcBorders')
            for e in ['top','bottom','left','right']:
                b=OxmlElement('w:'+e);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'CED9D1');borders.append(b)
            pr.append(borders)
            for a in c.paragraphs:
                a.paragraph_format.space_after=Pt(4);a.paragraph_format.line_spacing=1.15
                for r in a.runs:r.font.size=Pt(10);r.bold=ri==0
    return t
doc.add_paragraph('放风筝\n教学流程与参考答案','Title')
p('中二高级华文 · 单元一 · 巩固课文 · 2—3课时',True)
p('以14段课文、课本旁注及《巩固练习》为主线。一次处理一段，问题随段落推进；读完全篇后，完成扩写、人物比较和开放回应。三张原漫画保持画面、文字与格序。')
h('核心目标',2)
for i,o in enumerate(L['objectives'],1):p(f'{i}. {o}')
h('三课时完整安排',2)
rows=[]
for n,row in enumerate(M['flow3'],1):
    nums=[i+1 for i,s in enumerate(S) if s['lesson']==n];rows.append(row+[f'{min(nums)}—{max(nums)}页'])
table(['课次','课堂流程','自主输入','课件'],rows,[.6,3.35,2.0,.85])
p('每课50分钟，共150分钟。课内不显示倒计时。课件为46页，原生文字至少30pt；漫画保留原有字样。第一课用第3段短答，第二课用第12段短答与扩写，第三课用人物比较与开放回应。')
new('文本结构、困难与图文对应')
p('本文通过一次晚归冲突及海边放风筝，写“我”从感到委屈、抗拒管教，到理解父母爱与引导的过程。第一人称叙述使读者看见“我”当时的想法，不能把每一句内心评价当作客观定论。')
table(['段落','内容与功能','漫画格号'],[[f'{i+1}',idea,'、'.join(map(str,L['panels'][i]))] for i,idea in enumerate(L['paragraphIdeas'])],[.55,5.1,1.15])
p('第14格包含“我”的反应与爸爸赶紧去叫妈妈，横跨第8、9段；两段都显示同一完整格。学生页清楚标当前段号，教师提醒漫画包含紧接着的一小段行动。其他格按对应段落呈现，不重画或改台词。')
new('教学次序与文本边界')
h('学生困难与支架',2)
p('词句：战战兢兢、举白旗、五味杂陈、如梦初醒，须联系语境。阅读：分清“当时知道什么”与“后来才知道什么”。分析：把行动与特点、语言与用意连接。表达：先示范取证，再共构解释，最后独立人物比较与开放回应。')
table(['任务组','段落或范围','任务及课堂身份'],[
['1','第1—3段','先存疑；教师示范害怕的原因；共同比较父母；学生解释描写作用。'],
['2','第4—9段','点选词义，按短段顺序读到第8段后解释别扭。'],
['3','第10段','结合已读第4段，点选、取证解释五味杂陈。'],
['4','第11段','示例对照比喻与拟人，再口头分析课文。'],
['5','第12段','学生短答：明说的原因与合理推断分开。'],
['6','第13段','共同解释爸爸借风筝说话的用意。'],
['7','第14段','点选词义，口头回应比喻的道理与限制。'],
['8','全文后','先看课本扩写例句，独立扩写一句。'],
['9','全文后','自己概括爸爸、妈妈的特点，用原句解释。'],
['10','全文后','独立开放回应；自检、教师核对、订正。']],[.55,1.05,5.2])
p('10组任务拆成14个简短问题页，不把同类知识点跨段提前合并。第5—9段仍逐段读，未另加机械测试；教材问题全部保留。')
h('边界与讲评',2);p(M['boundary'])
p('课本将表情、神色与汗珠归为外貌描写，沿用教材口径，可以补充神态是更细的观察。第11段不设比喻／拟人二选一。第12段“改善关系”为教参提示的推断，明说的核心原因是帮助考试压力大的孩子解压。')
p('以下为课堂参考答案及建议意义点，不冒充官方试卷答案。口头题用于讲解，不计自动分；主观题由教师按意义核对。各独立意义点分别评估，不因漏答一部分而机械清零其他理解。')
for i,t in enumerate(L['tasks'],1):
    new(f'{i:02d}　{t["title"]}')
    p(t['source']+'　｜　'+t['role']+'　｜　能力：'+t['skill'])
    for q in t['questions']:p(q,True,True)
    h('定位与思考',2)
    p('范围：'+('、'.join(map(str,t['range'])) if len(t['range'])==1 else f'{t["range"][0]}—{t["range"][-1]}')+'段。'+t['guide'])
    keys=[w for w in t['keywords'] if any(w in q for q in t['questions'])]
    p(('题干关键词：'+'、'.join(keys)+'。' if keys else '')+'加工：'+t['processing'])
    for q in t['evidenceQuotes'][:3]:p(q,kai=True)
    h('参考答案',2);p(t['answer'],kai=True)
    h('答案点与接受范围',2)
    if t['points']:
        for pt,score in zip(t['points'],t['scores']):p(f'• {pt}（建议{score}点）')
    else:p('本页用于存疑与读文，不评分。')
    p(t['accept']+(' '+t['explain'] if t['explain']!=t['answer'] else ''))
    if t['id']=='mother-sun':p('前两点构成核心原因与取舍的答案；第三点是有依据的补充推断，未写不影响前两点。')
    h('易错点与分层提示',2);p(t['reject'])
    p('先提示：'+t['hints'][0]);p('再定位：'+t['hints'][1])
    for q,a in t.get('oral',[]):p('追问：'+q+'　参考：'+a)
    for alt in t.get('alternatives',[]):p('可接受方向：'+alt)
    p('证据边界：'+t['boundary'])
    # Compact misconception explanations for objective subquestions.
    for c in t.get('checks',[]):
        p('点选核对：'+c['question']+'　答案'+ 'ABCD'[c['correct']]+'。'+c['why'][c['correct']])
new('两课时流程与课堂操作')
table(['课次','课堂流程','自主输入'],M['flow2'],[.6,4.1,2.1])
p('每课50分钟，共100分钟。两课时版通过缩短重复汇报、第12段口头共构及扩写口答控制输入量；原题均保留。若学生取证、人物概括尚不熟练，使用三课时完整安排。')
h('互动页与学习记录',2)
for s in ['教师按段带读；学生读文页一次一段，原文左、对应原漫画右。问题页左当前段落、右本段题目。需要前文时用下拉菜单回看，不提前堆后文。','取证：点原句，再点对应框的“加入原句”。可累加多句、复制粘贴与编辑。第一次错误先给方向提示，5秒后可再试；第二次给更具体定位，再允许看讲解。','主观题保留首次作答，再展示参考、意义点与易错点供订正。系统只检查是否填写，不以字数、关键词判定答案质量。开放题必须写理由、例子与回扣。','“检查”页可填写姓名、班级和学号／组别，下载JSON、逐题CSV或文字记录。数据只存当前设备浏览器，不会自动发送教师或上传Google Sheet。','词语本汇总同域同浏览器个人收藏，按年级、单元、课文筛选。本课41项语境词库；未匹配的词可先建卡，清楚标“释义待补充”。闪卡先想再翻面，熟悉状态是自评。','重置本课保留词语本。清浏览器会丢本地记录，换设备前导出；共享设备应使用各自浏览器资料或分别备份。离线与在线记录不会自动互通。','教师页公开含答案，发给学生的是reading.html直达链接。在线投影与PPT同页序，问题先出现，讲评后揭。PPT文字和图形可编辑，漫画为原图。']:p('• '+s)
h('资源与来源',2)
p('教材：3张课文截图，共14段，改编自羊牧《放风筝》；1张《巩固练习》截图；3张教参截图。漫画：用户提供3张、共30格。教参所列课后作业页码保留为来源信息，因未提供作业页而未另造练习或布置新作业。')
p('学生：huawentong-home.lyqlym2015.workers.dev/du/textbook/2hcl/unit1/fang-feng-zheng/reading.html\n教师：同目录 teacher.html。网页与记录经过脚本及DOM验证；没有真实iPad或原生PowerPoint实测。')
doc.save(D/'downloads/teacher-guide.docx')
print(D/'downloads/teacher-guide.docx')
