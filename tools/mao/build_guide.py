from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[2];D=ROOT/'du/textbook/2hcl/unit2/mao';HERE=Path(__file__).parent
L=json.loads((HERE/'lesson.json').read_text());M=json.loads((HERE/'teaching_meta.json').read_text());S=json.loads((D/'slides.json').read_text())
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for b in doc.styles.element.xpath('.//w:pBdr'):b.getparent().remove(b)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Noto Serif CJK SC';st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=Pt(16);doc.styles['Normal'].paragraph_format.space_after=Pt(5)
for name,size in [('Title',25),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:doc.styles[name].font.size=Pt(size)
for name in ['Heading 1','Heading 2','Heading 3']:doc.styles[name].paragraph_format.space_before=Pt(9);doc.styles[name].paragraph_format.space_after=Pt(5)
footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('华文通　猫　教师版　');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(9)
def p(s,bold=False,kai=False):
 a=doc.add_paragraph();r=a.add_run(str(s));r.bold=bold
 if kai:r.font.name='LXGW WenKai TC';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'LXGW WenKai TC')
 return a
def h(s,level=1):doc.add_heading(s,level)
def new(s):doc.add_page_break();h(s)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,v in enumerate(headers):t.rows[0].cells[i].text=str(v)
 for vals in rows:
  cells=t.add_row().cells
  for i,v in enumerate(vals):cells[i].text=str(v)
 head=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(head)
 for ri,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','bottom','left','right']:
    b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
   pr.append(borders);margin=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margin.append(e)
   pr.append(margin)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE8EE');pr.append(sh)
   for a in c.paragraphs:
    a.paragraph_format.space_after=Pt(3);a.paragraph_format.line_spacing=1.15
    for r in a.runs:r.font.size=Pt(10.5);r.bold=ri==0
 return t
doc.add_paragraph('猫 教学流程与参考答案','Title')
p('中二高级华文　单元二　核心课文　3课时',True)
p('依据本次提供的12段课文、教材第37—43页与3张教参截图。教学重点是用外形、动作、声音的原文证据理解猫的形象，分清事实与误判，解释作者的愧疚。4张原漫画完整保留。')
h('核心目标',2)
for i,o in enumerate(L['objectives'],1):p(f'{i}. {o}')
h('课时安排',2)
rows=[]
for n,row in enumerate(M['flow'],1):
 nums=[i+1 for i,s in enumerate(S) if s['lesson']==n];rows.append(row+[f'{min(nums)}—{max(nums)}页'])
table(['课次','时间与内容','自主书写','课件'],rows,[.55,3.3,2.35,.9])
p('每课按50分钟，共150分钟。课件53页，原生文字至少30pt。各段问题按顺序推进，10组内容拆为17个问题页。第一课、第二课各1道重点书写；第三课2道。课内不设倒计时。')
new('文本结构与图文对应')
p('第一人称叙述让读者看见“我”的判断及反省。第1段提出愧疚；第2—4段交代印象；第5—9段写怀疑、定罪和责打；第10—11段发现误判并自责；第12段回到长久愧疚，与开头呼应。')
table(['段','内容与功能','漫画格号'],[[i+1,L['paragraphIdeas'][i],'、'.join(map(str,L['panels'][i]))] for i in range(12)],[.45,5.35,1.3])
p(M['comicNote'])
new('证据边界与教学支架')
p(M['boundary'])
h('学生困难与解决方法',2)
table(['困难','本课支架'],[['词句不懂','词语预习与课内点词，理解枕、仰、凝望、畏罪潜逃、衔、谴责、估计。已学词义不重复长讲。'],['把人物判断当事实','先比较“鸟死了”和“一定是猫”；再标出似乎、仿佛、好像、以为、猜想。'],['只报术语，不解释作用','先由教师示范第2段句子作用，再在第4、9段独立解释。第12段说明首尾呼应的具体内容。'],['只会抄句，不会概括','外形、动作、声音分别取证，再用自己的话概括形象与情感。'],['开放题只有口号','保留看法、理由、例子、回扣，明确追问为什么；不强迫学生选择宽恕。']],[1.65,5.45])
h('课堂身份与评估',2)
p('我做：第1段读出疑问，第2段示范句子作用。我们做：顺段取证，分析家人为什么误判，以及“我”自责的原因。你做：第4段反复强调的作用，第9段动作作用，以及最后的形象情感统整与角色回应。')
p('以下为教师课堂参考，不是官方试卷评分。每个意义点建议1分，允许等义表达；各点独立给分。缺少一项不清零其他独立理解。作用题只抄句子最多证明定位，须解释句子说明什么。网页仅对客观取证和点选给反馈，主观题待教师核对。')
p('第2、5、8、9段的后文关系随阅读逐步补全，不提前把未读段落堆到题目左栏。预习三的印象表按第2、3、4段分别完成；读完后再统整。')

def task(t,i):
 h(f'{i:02d} {t["title"]}',2)
 nums=[j+1 for j,s in enumerate(S) if s.get('task')==t['id']]
 p(t['source']+'　课堂：'+t['role']+'　课件：'+('、'.join(map(str,nums)) if nums else '见相关段落')+'页')
 for q in t['questions']:p(q,True,True)
 p('参考：'+t.get('fullAnswer',t['answer']),kai=True)
 if 'fullAnswer' in t:p('逐段讲评时先用：'+t['answer'])
 p('答案点（每点建议1分）：'+'；'.join(t['points'])+'。')
 p('证据：'+'；'.join(t['evidenceQuotes']),kai=True)
 p('思考：'+t['processing']+'。'+t['accept'])
 p('易错：'+t['reject'])
 p('提示1：'+t['hints'][0]+'　提示2：'+t['hints'][1])
 if t['mode']=='objective' and t.get('checks'):
  p('点选答案：'+'；'.join(f'第{j+1}项：'+'ABCD'[c['correct']] for j,c in enumerate(t['checks']))+'。干扰项分别检查时间倒置、人物对象、把猜想当证据等误读。')
 if t.get('oral'):p('口头追问：'+'；'.join(q+' '+a for q,a in t['oral']))
 if t.get('alternatives'):p('接受范围：'+'；'.join(t['alternatives']))
 if t['boundary']!='判断限于当前读到的课文信息。':p('边界：'+t['boundary'])

# Two related questions share a page only in this teacher reference, not on student screens.
groups=[[0,1],[2,3],[4,5],[6],[7,8],[9,10],[11,12],[13,14],[15],[16]]
for g in groups:
 new('问题讲评 '+str(g[0]+1)+(' 至 '+str(g[-1]+1) if len(g)>1 else ''))
 for i in g:task(L['tasks'][i],i+1)

new('教材表格与词语核对')
h('预习一与预习二',2)
table(['词语','本课语境义'],[['枕','把头放在“我”的脚上，像枕着枕头。'],['仰','抬头向上看。'],['畏罪潜逃','因害怕受罚而偷偷逃走；文中是家人的误判。'],['估计','根据情形推测或判断，不表示已经确定。'],['光泽','物体表面的亮光。'],['凝望','专注、长时间地望着。'],['惩罚','因对方被认为有错而给予处罚。'],['衔','用嘴叼着。']],[1.2,5.9])
h('预习三 印象表',2)
table(['时间','“我”对猫的印象'],[['刚收养时','毛色花白、没有光泽、瘦小、不好看；不活泼，一副忧郁的样子。'],['过了一段时间后','变胖仍不活泼；懒懒打哈欠；烧掉几块毛，更难看。'],['春天来了后','更胖却仍忧郁；不见捉鼠；终日懒洋洋、无精打采。']],[1.3,5.8])
h('阅读放大镜二 事情与反应',2)
table(['事情','猫的描写','人的感受或反应'],[['猫被收养','不好看、不活泼、忧郁。','大家都不太喜欢它。'],['家里买了鸟','跳到桌上，仰头凝望鸟笼；被赶走后又回来。','妻子提醒张妈留心猫，张妈把猫赶走。'],['死了一只鸟','露台晒太阳，神态安详，嘴里好像在吃什么。','“我”怒气冲天，用木棒狠狠地打猫。']],[1.1,3.2,2.8])
new('课堂使用与材料来源')
h('课堂操作',2)
for s in ['向学生发送互动学习直达链接。每课先逐段读文，再做相应问题；漫画直接显示。长段自然滚动，不设阅读解锁。','取证时先点左侧原句，再点右侧“加入原句”。连续加入会累加；可换段、复制、粘贴、编辑。','学生先答再看参考。客观题第一次错给方向提示，5秒后重试；第二次错给定位提示，可再看讲解。','第三课完成形象与情感、开放回应后，教师按意义点核对，自检不当作自动评分。开放题允许不同感受，但必须解释为什么。','在“检查”页填写姓名、班级、学号，导出JSON、逐题CSV或文字。记录含首次答案、订正、提示与提交时间。口头题不计自动正确率。']:p(s)
h('词语本与离线使用',2)
p('个人词语本与已接入的华文通课文共用。学生自己收藏，按年级、单元、课文筛选。55项本课词库提供语境义与原句；未知词也能建卡，标释义待补。翻卡后可标还不熟或已记住；这不是掌握测验。')
p('学习记录在同一设备、同一浏览器本地保存，没有账号云同步、自动发送教师或Google Sheet上传。重置本课保留词语本。清浏览器、换设备前先导出。离线版与在线版记录不会自动互通。中二单元五、六旧课的迁移仍暂缓。')
p('整包包括独立学生HTML、可编辑PPT、教师DOCX、在线投影的离线版、53张投影页、4张原漫画。先解压再打开。建议安装霞鹜文楷与思源宋体；原图文字不是可编辑文字。教师页公开含答案，不能把隐藏入口当成身份保护。')
h('资料范围',2)
p('课文：image(20261020-003003)、003017、003034、003114，教材第38—41页。预习与练习：002924、003152、003224，教材第37、42、43页。教参：003318、003337、003351。漫画：猫-1至猫-4。仅使用本课内容，未添加新短文或延迟小测。')
p('作者署名按教材：本文改编自郑振铎的《猫》。漫画为教师本次提供的既有资料，原样复用。所有答案为依据教材与教参整理的课堂参考，开放题示范不是唯一答案。')
doc.save(D/'downloads/teacher-guide.docx');print('DOCX generated')
