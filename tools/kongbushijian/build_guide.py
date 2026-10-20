from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
R=Path(__file__).resolve().parents[2];W=R.parent;O=R/'du/textbook/2hcl/unit4/kongbushijian';L=json.loads((O/'lesson.json').read_text());B=json.loads((O/'teaching-brief.json').read_text());T=L['tasks']
D=Document();sec=D.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=D.styles[name];st.font.name='Noto Serif CJK SC';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt({'Normal':11,'Title':24,'Subtitle':12,'Heading 1':16,'Heading 2':13,'Heading 3':12}[name]);st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=Pt(18 if name=='Normal' else 24)
D.styles['Title'].paragraph_format.space_after=Pt(14)
def p(s='',bold=False):
 a=D.add_paragraph();a.add_run(s).bold=bold;return a
def lead(a,b):
 q=D.add_paragraph();q.add_run(a+'：').bold=True;q.add_run(b);return q
def h(s,level=1):D.add_heading(s,level)
def page():D.add_page_break()
def table(headers,rows,widths):
 tb=D.add_table(rows=1,cols=len(headers));tb.autofit=False
 for i,w in enumerate(widths):tb.columns[i].width=Inches(w)
 for i,a in enumerate(headers):tb.rows[0].cells[i].text=str(a)
 for row in rows:
  for i,a in enumerate(row):tb.add_row() if False else None
  cs=tb.add_row().cells
  for i,a in enumerate(row):cs[i].text=str(a)
 for ri,row in enumerate(tb.rows):
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
   for edge in ['top','bottom','left','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');bd.append(x)
   pr.append(bd);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE7EB'if ri==0 else'FFFFFF');pr.append(sh);mar=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:w'),'90');x.set(qn('w:type'),'dxa');mar.append(x)
   pr.append(mar)
   for q in c.paragraphs:
    q.paragraph_format.line_spacing=Pt(16);q.paragraph_format.space_after=Pt(2)
    for r in q.runs:r.font.size=Pt(10.5);r.bold=ri==0
 tb.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'));return tb

def item(i):
 t=T[i];h(t['title'],2)
 for q in t['questions']:p(q,True)
 lead('来源与课堂身份',t['source']+'；'+t['role'])
 lead('参考答案',t['answer']);lead('证据与思考',t['explain']);lead('提示','；'.join(t['hints']));
 if t['reject']:lead('不接受',t['reject'])
 if t['mode']=='subjective':lead('评分建议',t['scoreGuide'])
 for q,a in t.get('oral',[]):lead(q,a)
D.add_paragraph('恐怖事件',style='Title');D.add_paragraph('中二高级华文 单元四 核心课文 教师流程与参考答案',style='Subtitle')
p('以“交流怎样改变关系”为主线，按19个自然段推进，保留五张原漫画。旧练习的有效训练已整合为10组、19个随段步骤；纠正了双关、互相道歉及无根据心理推测等解释。三课时方案循序展开，两课时方案精简汇报和书写量。')
h('核心目标')
for i,g in enumerate(B['goals'],1):p(str(i)+'　'+g)
h('三课时安排')
for r in B['schedule']:lead(r[0]+' '+r[1],r[2])
lead('课件分课','46页。第一课第1—18页，第二课第19—35页，第三课第36—46页。问题与答案分开呈现。')
page();h('两课时安排与教学取舍')
p(B['twoLesson'])
lead('课堂思考示范','我做：第2段把行为与影响组成原因。我们做：第6段写法—具体气氛—人物关系；第19段关键句与情节。你做：第18段解释两个词的语境义、感叹句改写、必做开放回应。')
lead('避免无依据推断','“不敢反驳”只支持当时不敢把话说完；未明写服气、委屈或认错。其他室友“不愿多说”不能证明她们内向。原文“我”不等同于已核实的作者本人。')
h('材料与图画边界')
for s in B['sourceBoundary'][:3]:p(s)
lead('评改原则','参考答案及分值是依据课文和教参整理的课堂建议，不冒充官方评分细则。接受有依据的等义表达，按独立意义单位给分。开放题与作用题不凭字数或关键词命中评分。')
page();h('段落 图画与问题对应')
table(['段落','内容重点','原图格号'],[[i+1,L['paragraphIdeas'][i],'、'.join(map(str,L['panels'][i]))]for i in range(19)],[.55,5.1,1.35])
lead('阅读呈现','逐段左文右图，图上写中文段号，图下积累词语。问题页左段落右问题；先按段落推进，最后才统整全文。')
# Compact teaching spreads, with extra space reserved for major written answers.
for title,idxs in [('第一课 话语与作息冲突',[0,1]),('第一课 提醒与反应',[2,3]),('第一课 关系与夸张',[4,5]),('第一课 开门与眼神',[6,7]),('第一课 动作取证',[8]),('第二课 声音与动作',[9,10]),('第二课 共同点与转变',[11,12]),('第二课 真相与误会',[13,14]),('第二课 关键句与主题',[15]),('第三课 统整与感叹句',[16,17]),('第三课 必做开放回应',[18])]:
 page();h(title)
 for i in idxs:item(i)
 if idxs==[8]:
  lead('多句证据示范','先加入“我鼓起勇气拉开门，小心地探出身望了望，走廊空荡荡的，什么也没有。”再加入“过后，我们赶紧熄灯上床。”两句保留后，解释“小心”“赶紧”怎样反映恐惧。第9、12、13段的证据随段处理，不提前堆在同一页。')
 if idxs==[15]:
  lead('抒情句的可接受范围','第1段也直接表达“我”的情感，并非只有第18段才有抒情。课本要求结合主题时，优先选择第18段的“惊悚而美丽”，便于联结第19段议论与第17段关系变化。不能用“人物对话一定不是抒情句”作机械排除规则。')
 if idxs==[18]:
  for s in T[18]['alternatives']:p(s)
page();h('旧练习整合与删除说明')
p('用户提供的两个网址相同。仓库中只找到这一份24题旧页面，主页将它同时列入“单元四课文阅读理解”和“课前预习”。新版直接替换该网址的旧内容，只保留“单元四 核心课文 恐怖事件”入口。Git历史保留原版本，不再另挂旧练习。')
table(['旧题','处理','理由'],[
['Q1—5','整合至全文统整与关键句','先逐段读懂，再归纳，修正结构和关键句判法'],['Q6—8、Q12','随第2—4段修订','保留习惯、反应；删除内向、委屈、服气的确定性猜测'],['Q9—11','拆回对应段落','恐惧的眼神、动作和声音按顺序取证'],['Q13—16','整合第6段','夸张与动词效果保留，删荒谬干扰及神秘力量等添写'],['Q17—20','随第5、17、18段修订','冤家不等于亲近；讨厌鬼只指蒋妍；分别解释惊悚美丽'],['Q21','删除','无法确定“我”与真实作者身份关系'],['Q22—24','整合主题与必做开放回应','保留理由、生活经历和回扣，取消关键词自动判分']],[1.0,2.0,4.0])
p('删除侦探闯关、连击和经验值包装，采用已确认的课文学习流程。旧版state只存页面内存，未发现可迁移的持久答案；新版采用独立课文标识s2u4l2，不套用其他课答案。')
page();h('课堂操作与记录')
for label,txt in [('学生入口','reading.html，读文、三课任务、检查与词语本。课文主页保留原网址。'),('取证','点左边原句，再点“加入原句”；可连续加入、复制粘贴、换行和编辑，重复同句不会重复添加。'),('反馈','客观题第一次答错先给提示，5秒后可重试；再错可看讲解。主观题先作答再展示参考，首次答案不被订正覆盖。'),('必做开放题','须填写立场理由、真实经历和回扣。可不同意或部分赞同，重点看解释与例证。不用“已讨论”代替表达。'),('个人词语本','本课55项语境词库；收藏学生自己需要的词，不自动加入全部词表。未知词仍建卡，标释义待补充，补充前不能标已记住。'),('记录边界','当前设备、同一浏览器保存。检查页导出JSON、CSV和文字记录；没有账号云同步或教师自动收件。清除浏览器前先导出。重置本课保留词语本。'),('跨课文收藏','同域已接入课文共享收藏，按年级、单元、课文筛选，保留各出处、释义与熟悉状态。站外旧课文不能自动汇总。'),('课件','46页可编辑PPT。文字及两页统整表格可编辑，漫画为保留的原图。字体为霞鹜文楷与思源宋体，未安装时可能替换。在线投影为同版渲染图片。'),('离线','独立HTML可直接打开；整包ZIP先解压，再打开“开始使用.html”。外部预习页需要网络，朗读依赖设备语音包。'),('公开范围','教师页为公开资源页，分开入口不等于登录保护。'),('核验范围','完成内容、资源、PPT和DOCX渲染及DOM交互检查；没有实体iPad或原生PowerPoint实机测试。')]:lead(label,txt)
foot=sec.footer.paragraphs[0];foot.alignment=2;foot.add_run('恐怖事件  ').font.size=Pt(9);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
for root in [D.element,D.styles.element]:
 for border in root.xpath('.//w:pBdr'):border.getparent().remove(border)
D.save(O/'downloads/teacher-guide.docx');print('Teacher guide ready')
