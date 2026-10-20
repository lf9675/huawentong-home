from pathlib import Path
import os
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
W=Path(os.environ.get('XQ_BUILD_DIR','/tmp/huawentong-xinqing-build'));W.mkdir(parents=True,exist_ok=True);O=Path(__file__).resolve().parents[2]/'du/textbook/2hcl/unit4/xin-qing-shou-zhang';L=json.loads((O/'lesson.json').read_text());B=json.loads((O/'teaching-brief.json').read_text());T=L['tasks'];D=Document()
sec=D.sections[0];sec.page_height=Inches(11.7);sec.page_width=Inches(8.3);sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.72)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
 s=D.styles[name];s.font.name='Noto Serif CJK SC';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');s.font.size=Pt({'Normal':11.5,'Title':24,'Heading 1':19,'Heading 2':14,'Heading 3':12}[name]);s.font.color.rgb=RGBColor.from_string('000000'if name=='Title'else'243D37');s.paragraph_format.space_after=Pt(6)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
 D.styles[name].paragraph_format.line_spacing=Pt(18 if name=='Normal' else 24)
 D.styles[name].paragraph_format.space_before=Pt(0 if name in ['Normal','Title'] else 8)
 D.styles[name].paragraph_format.space_after=Pt(5)
D.core_properties.title='心情手账 教师流程与参考答案';D.core_properties.author='华文通'
def p(s='',style=None):return D.add_paragraph(s,style)
def h(s,level=1):D.add_heading(s,level)
def new(s):D.add_page_break();h(s)
def table(headers,rows,widths=None):
 t=D.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
 if widths:
  for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for c,s in zip(t.rows[0].cells,headers):c.text=s;sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF1ED');c._tc.get_or_add_tcPr().append(sh)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for c,s in zip(cells,row):c.text=str(s)
 for row in t.rows:
  no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
  for i,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[i])
   for para in c.paragraphs:para.paragraph_format.space_after=Pt(4);para.paragraph_format.line_spacing=Pt(16)
 return t
def tasktext(t):
 p('来源：'+t['source']+'。课堂身份：'+t['role']+'。')
 for q in t['questions']:p(q)
 h('参考答案',2);p(t['answer']);h('依据与讲解',2);p(t['explain'])
 h('提示与常见误读',2);p('第一次提示：'+t['hints'][0]);p('第二次提示：'+t['hints'][1]);p('不接受：'+t['reject']);
 if t['id']!='character':p('核对方式：'+t['scoreGuide'])
p('心情手账 教师流程与参考答案','Title');p('中二高级华文　单元四　生活空间　1课时')
p('本课以四则心情手账为材料，帮助学生读懂烦恼与应对打算，再用文本证据分析人物、提出有理由的看法。按50分钟安排，保留全部课本练习；点选减轻输入，人物分析与开放回应分别完成一次自主书写。')
h('教学目标',2)
for i,g in enumerate(B['goals'],1):p(str(i)+'. '+g)
h('课堂流程',2);table(['环节','时间','教学与产出'],B['schedule'],[1.0,.65,4.9])
p('课件按逐则顺序推进，题目与答案分开。学生先答后揭示；班级需要时减少重复汇报，不删课本表格与必做开放题。')
new('文本与教学依据')
p('四则手账各自记录一个烦恼，并写出反应或打算。它们按日期排列，但不是同一事件从起因到结果的四阶段。第一人称“我”的感受和评价，需要与具体事件及未确定情况分开。')
table(['日期与漫画','要点与原文证据','阅读加工'],[
['7月1日\n第1—6格','“可是，我又无法抗拒零食的诱惑。”“我要做一个有毅力的人……”','连接矛盾两面，分清愿望与已实现结果'],
['7月2日\n第7—12格','“好像越来越疏远”“好像是睡眠不足”“我明天要……”','区分观察、推测和计划'],
['7月3日\n第13—18格','“却把……全让给我”“既然……那……”','解释行动与情绪变化的联系'],
['7月4日\n第19—24格','“我想选修……”“他希望……”“但我还是想……”','概括两个方向的差别及个人选择']], [.9,3.5,2.0])
h('学生常见困难与支架',2)
p('词语：自卑、毅力、闪烁其词、怄气等。直接在原文中理解语境，课本词语题随则完成。篇章：每则只整理“烦恼、打算”两项，不混入别则。分析：用实际行为证明特点，不能用愿望证明已经具备的品质。表达：看法后必须解释为什么，再用细节说明。')
h('教材与漫画边界',2)
for x in B['sourceBoundary']:p(x)
for n in range(1,5):
 new('第'+'一二三四'[n-1]+'则 7月'+str(n)+'日')
 h('课文原文',2);para=p(L['paragraphs'][n-1])
 for run in para.runs:run.font.name='LXGW WenKai TC';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'LXGW WenKai TC')
 vt=next((t for t in T if t['id']=='words-july'+str(n)),None)
 if vt:
  h('课本词语题',2);p(vt['questions'][0]);p('　'.join(b['before']for b in vt['blanks']));p('参考：'+vt['answer']);p(vt['explain'])
 h('课本表格',2);t=next(t for t in T if t['id']=='table-july'+str(n));tasktext(t)
 if t.get('oral'):
  for q,a in t['oral']:p('口头追问：'+q);p('参考：'+a)
 if n==1:p('教师提醒：理解主人的自卑感受，同时避免用身材评价同学的价值。无需分享个人体重，也不将本课变成减重指导。')
 if n==2:
  h('证据能支持到哪里',2);table(['原文信息','可以判断','不能据此断定'],[['多次询问，又打算谈一谈','主人关心并愿意了解朋友','已经知道她的问题、已经帮助她解决'],['好像是睡眠不足','主人提出一种可能解释','真实原因必然是熬夜或游戏']],[2.0,2.0,2.4])
new('人物分析 自主作答')
t=next(t for t in T if t['id']=='character');tasktext(t)
h('可接受的不同分析',2)
for x in t['alternatives']:p(x)
h('意义单位评分建议',2)
table(['项目','分值','判断依据'],[['身份与阶段','1','即将进入中三的中学生，或有依据的青少年阶段；不要求精确岁数。'],['特点','1','特点与原文相符，不用漫画给出的身份。'],['依据','1','具体且有关联的实际行为或打算；不得夸大为已成功。'],['解释','2','说明行为体现的态度或考虑，建立联系；解释部分成立可给1分。']],[1.3,.5,4.6])
p('各意义单位独立给分，不因漏写身份而清零后面合理的人物分析。若证据不支持所写特点，该证据不能为该特点得分，但其他独立正确内容保留。取证框支持回看四则、连续加入多句、编辑和粘贴。')
new('开放回应 必做')
t=T[-1];p('来源：'+t['source']);p(t['questions'][0]);p('课堂身份：学生独立。先不显示示范，给9分钟用于写答、比较和订正。可赞成、部分赞成或提出不同看法，必须解释理由。')
table(['步骤','提问','判断'],[['看法','我怎么看？建议什么？','明确选定一则，观点相关。'],['理由','为什么？','解释需要、影响或考虑，避免循环理由。'],['例子','哪一细节或做法能说明？','结合原文或具体可行的情境。'],['回扣','怎样支持看法？','连接做法、作用与看法。']],[.7,1.7,4.0])
h('7月2日参考示范',2);p(t['answer'])
h('其他三则的合理方向',2)
for x in t['alternatives']:p(x)
h('教师评分与反馈',2);p(t['scoreGuide']);p('提示一：'+t['hints'][0]);p('提示二：'+t['hints'][1]);p('常见问题：'+t['reject']);p('示范不要求学生照抄。建议是否合理，结合所选日记与说明判断；提及沟通不自动算正确。')
new('使用与核对')
h('资源与课堂使用',2)
p('课文主页：开始／继续学习、四则直达、词语预习、图文回读与教师资料。互动页：原文在左，漫画或问题在右，窄屏上下排列；每则原漫画六格按原顺序展示。先读本则，再点“本则问题”，按日期推进。')
p('教师页：在线投影、可编辑PPT、教师DOCX、独立离线HTML及整套ZIP。投影与PPT同版，题目页与答案页分开，左右方向键翻页，教师提示默认收起。')
h('题目覆盖',2);table(['附件原题','实现位置'],[['正文“用几个词语概括一下你这一周的心情”','导入口头完成，不新增文字任务'],['阅读放大镜一（1）（2）（3）（4）','第1、4、6任务，共四个释义找词'],['阅读放大镜二表格四行','第2、3、5、7任务，按四则逐行点选'],['教参2.1人物推测','第8任务，自主特点、依据与解释'],['阅读放大镜三','第9任务，必做开放回应']],[3.0,3.4])
h('记录与设备范围',2)
p('九个任务中七个为客观任务：三组词语、四行表格。首次正确率按“整项任务是否全部正确”统计，主观两题不计入自动正确率。教师建议人物题5分、开放题5分，二者合计10分，仅供人工形成性评价；不与网页正确率混算。')
p('保存姓名、班级、学号／组别、首次答案、每次提交、时间、提示使用、订正和自检。检查页可导出JSON完整备份、逐题CSV与文字记录。本课用独立存储标识，重置不删除个人词语本。未接入账户同步或教师自动收件。')
p('个人词语本汇总已接入同域课文的个人收藏，按年级、单元、课文筛选。预设29项语境词并非强制全部收藏；自输未知词仍可建卡，显示“释义待补充”。导出后可另机导入，清理浏览器前先备份。设备朗读不是人工录音。')
p('离线HTML包含本课原图、字体和脚本。整套ZIP全部解压后打开“开始使用.html”。平台词语预习需网络，朗读取决于设备语音。教师资料是公开网页，不视为身份保护。')
h('当堂检查',2);p('只检查已有回答：有没有把打算当结果？特点有没有证据？看法有没有为什么？不增加新短文、课后一周小测或打卡作业。')
# Page numbers, without repeating oversized headers.
foot=sec.footer.paragraphs[0];foot.alignment=2
run=foot.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
for root in [D.element,D.styles.element]:
 for border in root.xpath('.//w:pBdr'):border.getparent().remove(border)
D.save(O/'downloads/teacher-guide.docx')
print('Teacher guide written with 9 planned pages.')
