from pathlib import Path
import json,sys
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

L=json.loads(Path(sys.argv[1]).read_text());ROOT=Path(__file__).resolve().parents[2];D=ROOT/'du/textbook/2hcl/unit1/ba-ai-shuo-chu-lai'
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.68);sec.left_margin=sec.right_margin=Inches(.72)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name];st.font.name='Noto Serif CJK SC';st.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=Pt(15)
doc.styles['Normal'].paragraph_format.space_after=Pt(4)
doc.styles['Title'].font.size=Pt(25);doc.styles['Heading 1'].font.size=Pt(18);doc.styles['Heading 2'].font.size=Pt(14);doc.styles['Heading 3'].font.size=Pt(12)
for n in ['Heading 1','Heading 2','Heading 3']:doc.styles[n].paragraph_format.space_before=Pt(7);doc.styles[n].paragraph_format.space_after=Pt(4)
header=sec.header.paragraphs[0];header.text='华文通　中二高级华文　单元一';header.style='Normal';header.runs[0].font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('把爱说出来　教师版　');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def p(text,bold=False,kai=False):
    q=doc.add_paragraph();r=q.add_run(text);r.bold=bold
    if kai:r.font.name='LXGW WenKai TC';r._element.rPr.rFonts.set(qn('w:eastAsia'),'LXGW WenKai TC')
    return q
def h(text,level=1):
    para=doc.add_heading(text,level)
    if level==1:para.paragraph_format.page_break_before=True
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,x in enumerate(headers):t.rows[0].cells[i].text=x
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        c=t.add_row().cells
        for i,x in enumerate(row):c[i].text=str(x)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),'D9D9D9');b.append(el)
            pr.append(b);mar=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E6ECE8');pr.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.1
                for run in para.runs:run.font.size=Pt(10);run.bold=ri==0
    return t
doc.add_paragraph('把爱说出来教学流程与参考答案','Title')
p('中二高级华文　单元一　建议三课时　每课50分钟',True)
p('本课通过比较女儿、儿子和“我”表达爱的方式，引导学生由语言、行动和心理找证据，解释“我”开口前的困难，并理解结尾歌声所表现的爱的力量。配套优先保留课本原题，采用点选、取证与少量自主书写。')
h('核心目标',2)
for x in L['objectives']:p('• '+x)
h('三课时安排',2)
table(['课次','时间与课堂活动','课件及输入'],[
['第一课','导入3分钟；第1–8段阅读15分钟；任务1–3共15分钟；第9–11段及任务4共12分钟；订正与词语5分钟。','课件1–17页。任务3简答；其余点选、取证和口头解释。'],
['第二课','复习3分钟；第12–15段阅读12分钟；任务5用8分钟；任务6用10分钟；任务7用12分钟；订正5分钟。','课件18–30页。任务5与7书面作答。'],
['第三课','第16段及回读8分钟；任务8用12分钟；任务9讨论8分钟；任务10独立回应12分钟；自检订正与词语10分钟。','课件31–45页。任务9简记与任务10完整回应。']],[.65,4.2,2.15])
p('总计150分钟。预计时间用于教师掌握节奏，不显示倒计时，不给速度分，不公开个人排名。两课时压缩方法见末页。')
p('教材来源：本次上传课文4张、阅读放大镜及读后说2张、教参3张、原有漫画3张。课文共16段，改编自廖玉蕙《示爱》。',kai=False)
h('逐段内容与漫画对应')
p('逐段原文保存在网页content.js；以下为概括，不替代课文原句。漫画只作理解支架，答题证据须来自原文。三张原图保持原有画面与文字，仅用显示窗口定位相应完整格。')
rows=[]
for i,(idea,panels) in enumerate(zip(L['paragraphIdeas'],L['panels'])):
    taskids=[j+1 for j,t in enumerate(L['tasks']) if (len(t['range'])==2 and t['range'][0]<=i+1<=t['range'][1]) or (len(t['range'])!=2 and i+1 in t['range'])]
    rows.append([i+1,idea,'、'.join(map(str,panels)),'、'.join(map(str,taskids))])
table(['段','内容要点','原图格号','任务'],rows,[.4,4.25,1.15,1.2])
p('第5段对应第8格；这一格兼含卡片收藏信息。第6、7段分别对应第9、10格。第9段需完整展示第13–16格，避免只看母亲做家务而漏掉丧偶后的心情。')
h('文本分析与证据边界')
h('结构与主题',2)
p('第1–8段：用女儿与儿子对文章、新衣的不同反应，写出女儿的赞美及卡片使“我”温暖。第9–11段：“我”感激母亲，受女儿启发，决定表达爱。第12–15段：“我”害羞退缩，最终开口。第16段：母亲重新唱歌，突出表达爱的力量，呼应题目。')
h('学生容易卡住的地方',2)
p('词句：把“碰钉子、泼冷水、疙瘩”按字面理解；把“甜言蜜语”一概看成欺骗。结构：只记得女儿赞美妈妈，没看出“我”由受爱者转为表达者。推断：把儿子不赞美等同于不爱；把一次歌声扩大为永不悲伤。表达：只说“生动形象”“不习惯”，缺少对象、程度与证据。')
h('事实与推断',2)
p('明示：女儿赞叹、重读文章、写卡片；儿子拒听并否定妹妹；“我”从未向母亲表达，面对她害羞；母亲在“我”表达后唱歌。合理推断：母亲感受到爱与安慰，心情有所转变。无法确定：儿子心中爱的程度、特定文化如何形成“我”的性格、母亲此后的全部情绪。')
h('教参与本配套的表述差异',2)
p('教参以“子女对母亲的爱与生俱来”引导儿子可能也爱母亲。本配套保留“表达方式不同”的教学意图，但不以普遍断言替代文本证据，接受“可能爱，只是不习惯表达”及其他有据、有保留的看法。')
p('教参结尾答案写母亲“从父亲过世的阴影中走了出来”。本配套以“受到爱与安慰，心情有所好转”表达其核心含义；若学生写“走出阴影”，应确认他指文中此刻的转变，不能扩大为从此完全不再悲伤。')
p('课本阅读放大镜把女儿“双眼含笑、一脸欣赏”等列为外貌描写。课堂保留该教材分类，并说明更细可称神态描写。第10段以心理描写为主；“开车”不是整段主体。')
h('教学取舍',2)
p('不把《团圆饭》的线索与引号任务搬入本课。不加新短文、不加课后一周小测。课本“读后说”是必做开放回应，不以勾选“已讨论”代替理由与例证。课本完整人物题保留，由学生自主写特点。')
for i,t in enumerate(L['tasks']):
    h(f'任务{i+1}　{t["title"]}')
    p(f'第{t["lesson"]}课　{t["role"]}　来源：{t["source"]}',bold=True)
    for q in t['questions']:p(q,kai=True)
    if t.get('blanks'):
        for j,b in enumerate(t['blanks']):p(f'{j+1}. {b["before"]}________{b["after"]}',kai=True)
    h('审题与取证',2)
    p('题干关键词：'+'、'.join(t['keywords'])+'。')
    if t.get('evidence'):
        for e in t['evidence']:p(e['label']+'：第'+'、'.join(map(str,e['p']))+'段；注意“'+'”“'.join(e['terms'])+'”。')
    else:p('定位范围：第'+(('至'.join(map(str,t['range']))) if len(t['range'])==2 else '、'.join(map(str,t['range'])))+'段。')
    p('所需加工：'+{'objective':'定位并比较、辨析或整合信息；检查后说明为什么。','written':'筛选证据，概括意义并解释关系，组织独立书面答案。','discussion':'明确立场，解释适用条件，用课文支持判断。'}[t['mode']])
    h('参考答案',2);p(t['answer'],kai=True)
    h('意义得分点与接受范围',2)
    for point,score in zip(t['points'],t['scores']):p(f'• {point}（建议{score}分）')
    if len(t['scores'])!=len(t['points']):p('补充计分：惯用语每空1分，共3分。')
    p('以上分值为课堂讲评建议，并非教材公布的官方评分。按独立意义单位给分，接受等义表达。引用原文可作证据，但须筛选并回应题意；整段照抄不等于完成分析。')
    for alt in t.get('alternatives',[]):p('可接受方向：'+alt)
    p('不接受或需订正：'+t['reject'])
    h('分层提示与讲评',2)
    p('第一次：'+t['hints'][0]);p('第二次：'+t['hints'][1]);p('讲评：'+t['explain'])
    if t.get('oral'):
        h('保留的教参口头追问',2)
        for q,ans in t['oral']:p(q,bold=True);p(ans)
    if t['mode']!='objective':p('先保留首次回答，比较得分意义后在订正框补足。网页不按字数或关键词给主观题评分；教师覆核观点、证据及关系。')
h('两课时流程与使用说明')
h('两课时各50分钟',2)
p('第一课：导入3分钟；第1–8段图文与任务1–3共29分钟；第9–11段及任务4共13分钟；订正5分钟。保留任务3简答，缩短重复汇报。')
p('第二课：第12–16段阅读10分钟；任务5口头共构5分钟；任务6点选比较6分钟；任务7独立人物回答10分钟；任务8整合5分钟；任务9口头讨论4分钟；任务10独立回应7分钟；订正3分钟。合计50分钟。')
p('两课时版的任务5、9通过投影口头完成，学生页可继续访问但不以提交这些题作为完成全课的门槛。教师不把两题的共构答案当作独立掌握。人物比较和最后开放回应仍须独立作答。程度较弱、需要逐题书写的班级优先用三课时版。')
h('老师上课怎样用',2)
for x in ['从华文通首页依次进入“中二—课文学习—单元一—核心课文 把爱说出来”。','给学生发reading.html直达链接；老师打开teacher.html，再点“在线投影”。教师页公开含答案，不是受保护后台。','课件问题与讲解分开页；先让学生答，再翻到答案页。在线投影与45页可编辑PPT对应，课堂可随时跳页。','读文时左原文、右中文段号和对应漫画，下面收藏词语；任务中先点原句，再点证据框，iPad不必依赖拖放。','客观题首次错误给方向提示，锁定5秒后重试；第二次错误给更具体提示，再可查看讲解。主观题提交后对照、自检、订正。']:p('• '+x)
h('保存与下载',2)
p('本课独立存储键：huawentong-lesson-2hcl-u1-ba-ai-shuo-chu-lai-v1。首次答案、订正、每次尝试和时间、提示使用、阅读位置及填写的姓名班级学号保存在当前浏览器。检查页可导出JSON、CSV及文字。没有账号云同步，没有Google Sheet自动上传，不会自动发送老师。')
p('个人词语本沿用huawentong-personal-words-v1和huawentong-words-v1格式；有本课／全部筛选、编辑、导入导出与翻卡。未收录词可建卡，释义标为待补；未有释义不能标已记住。同词保留已有解释与熟悉状态，另列本课语境义。重置本课不删词语本。')
p('整套下载包含可编辑PPT、教师DOCX、原漫画、逐页投影图及独立互动HTML。离线HTML可直接打开；若浏览器限制file地址保存，使用记录导出备份。离线版与在线版的本地记录不会自动互通。')
p('验证范围：原文与附件逐段核对、脚本及DOM交互模拟、文件渲染、资源与下载校验。尚无真实iPad实测。正式使用前可在学校设备打开学生链接，确认中文显示、点句和输入。')
# Disable inherited East Asian grid snapping and decorative title borders.
for el in list(doc.element.iter(qn('w:docGrid'))):el.getparent().remove(el)
for el in list(doc.styles.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
for el in list(doc.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
for para in doc.paragraphs:
    pr=para._p.get_or_add_pPr();e=OxmlElement('w:snapToGrid');e.set(qn('w:val'),'0');pr.append(e)
dest=D/'downloads/teacher-guide.docx';doc.save(dest);print(dest)
