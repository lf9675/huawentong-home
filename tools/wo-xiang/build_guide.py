from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[2];D=ROOT/'du/textbook/2hcl/unit1/wo-xiang-dui-ni-shuo';HERE=Path(__file__).parent
L=json.loads((HERE/'lesson.json').read_text());M=json.loads((HERE/'teaching_meta.json').read_text());S=json.loads((D/'slides.json').read_text())
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.8);sec.footer_distance=Inches(.3);sec.left_margin=sec.right_margin=Inches(.7)
for b in doc.styles.element.xpath('.//w:pBdr'):b.getparent().remove(b)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Noto Serif CJK SC';st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=Pt(16);doc.styles['Normal'].paragraph_format.space_after=Pt(5)
for name,size in [('Title',25),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:doc.styles[name].font.size=Pt(size)
for name in ['Heading 1','Heading 2','Heading 3']:doc.styles[name].paragraph_format.space_before=Pt(9);doc.styles[name].paragraph_format.space_after=Pt(5)
footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('华文通　我想对你说　教师版　');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(9)
def p(s,bold=False,kai=False):
 a=doc.add_paragraph();r=a.add_run(str(s));r.bold=bold
 if kai:r.font.name='LXGW WenKai TC';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'LXGW WenKai TC')
 return a
def h(s,level=1):doc.add_heading(s,level)
def new(s):doc.add_page_break();h(s)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,v in enumerate(headers):t.rows[0].cells[i].text=str(v)
 for vals in rows:
  cells=t.add_row().cells
  for i,v in enumerate(vals):cells[i].text=str(v)
 head=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(head)
 for ri,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','bottom','left','right']:
    b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
   pr.append(borders);margin=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margin.append(e)
   pr.append(margin)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE8EE');pr.append(sh)
   for a in c.paragraphs:
    a.paragraph_format.space_after=Pt(3);a.paragraph_format.line_spacing=1.15
    for r in a.runs:r.font.size=Pt(10.5);r.bold=ri==0
 return t
doc.add_paragraph('我想对你说 教学流程与参考答案','Title')
p('中二高级华文　单元一　生活空间　1课时',True)
p('本课用三组求助与回应训练信息筛选、语境词义和有理由的表达。按消息顺序读文：国明先示范，美华共同整理，匿名同学独立填写表格，最后完成开放回应。以下为课堂参考答案，并非教材提供的官方评分标准。')
h('核心目标',2)
for i,o in enumerate(L['objectives'],1):p(f'{i}. {o}')
h('50分钟课堂流程',2)
table(['时间（分）','环节','课堂操作'],[[r[0].replace('分钟',''),r[1],r[2]] for r in M['flow']],[.85,1.45,4.8])
p('8组任务，28页可编辑PPT。第1—2页导入；3—8页国明；9—14页美华；15—20页匿名同学；21—24页词语；25—27页开放回应；28页检查。主要书面输入为匿名表格与开放回应两项。',True)
new('文本结构与阅读边界')
p(M['boundary'])
table(['教学段','消息与内容','漫画格'],[[str(i+1),L['paragraphLabels'][i]+'\n'+L['paragraphIdeas'][i],'、'.join(map(str,L['panels'][i]))] for i in range(6)],[.7,5.15,1.25])
p('三张原漫画共24格，原图未改。六段依次对应1—4、5—8、9—12、13—16、17—20、21—24格。第7—8、16、24格仅为建议情境；第15格保留“也许”，第22格保留“有些人的情况”。')
h('易混信息',2)
p('第一段：爸爸疲惫、不开心是背景；国明想安慰又怕打扰是他的困难。第三段：妈妈说美华叛逆，是人物评价，不能直接认定为事实。第四段：“如果没有”是条件，“说不定”是可能。第五段：“没有自由”是同学的感受；第六段：“有些人”不能换成匿名同学。')
h('语言困难与支架',2)
p('用“想……又怕……”找两难，用“我解释，她却……”理解委屈，用“不但……也……所以……”理清担心与限制的关系。词语点击释义按需使用；重点练习后才显示完整讲评。')
p('迁移能力：按对象筛选信息、区分原因与表现、保留范围和条件、以具体情境支持观点。不添加人物性格题、线索题或新短文。')
def task(t,i):
 h(str(i)+' '+t['title'],2)
 p(t['source']+'　'+t['role'])
 for q in t['questions']:p(q,True,True)
 p('参考：'+t['answer'],kai=True)
 p('依据：'+'；'.join(t['evidenceQuotes'][-1:]),kai=True)
 p('答案点：'+'；'.join(t['points'])+'。')
 p('思考：'+t['explain'])
 p('易错：'+t['reject'])
 p('提示1：'+t['hints'][0]+'　提示2：'+t['hints'][1])
for ids,title in [([0,1],'国明 教师示范'),([2,3],'美华 师生共构'),([4,5],'匿名同学 独立作答')]:
 new(title)
 for i in ids:task(L['tasks'][i],i+1)
new('课本表格与词语核对')
p('课本《阅读放大镜》二：根据课文内容，写出三位同学的烦恼及辅导员的建议。三行均保留，按对话出现位置完成。',True)
table(['同学','烦恼','辅导员的建议'],[['国明','爸爸疲惫、不开心；自己想安慰又怕打扰，不知怎样做。','直接表达关心，可写卡片打气或拥抱爸爸。'],['美华','做小组作业晚归，受妈妈责备，解释被当成顶嘴。','体谅妈妈因孩子晚归而担心，找机会好好沟通。'],['匿名','每天只准玩半小时游戏且被监督，觉得不自由、不被理解。','从爸爸角度思考，理解他担心学习及家人关系受影响。']],[.65,3.2,3.25])
p('接受保留原意的概括，不要求唯一措辞。国明、美华两行以口头共构为主，学生可按需在网页记录；匿名一行须自主填写。原句可证明定位，完整照抄辅导员的话若未筛选建议，不视为完成概括。')
h('课本词语题 六项',2)
table(['题号','课本释义','答案'],[[i+1,m,w] for i,(w,m) in enumerate(L['idioms'])],[.6,5.3,1.2])
p('“委屈”源于受到不应有的指责；“内疚”源于觉得自己有亏欠。第五项考窒息的基本义，人物留言中的“快要窒息了”形容压抑。六题用同一词语库点选，不额外要求长篇造句。')
new('开放回应与课堂评价')
t=L['tasks'][-1]
for q in t['questions']:p(q,True,True)
p('第一句为课本《阅读放大镜》三；第二句与3—4句表达支架为教师补充。此题必做。学生可选父母或辅导员，说出愿意分享的想法；也可明确使用假设情境，不要求公开私人经历。')
table(['环节','内容','教师核对'],[['想说的话','对象与希望','有没有说明想让对方理解什么。'],['理由','为什么要说','能否解释自己的需要或感受。'],['例子','具体情境或做法','情境是否支持理由，假设须标明。'],['回扣','怎样有帮助','能否解释做法与沟通之间的联系。']],[1.15,1.8,4.15])
p(t['answer'],kai=True)
p('可接受：向辅导员求助、向父母说明需要、提出有理由的不同意见或表达感谢。不能以是否同意父母决定得分。不接受：只写“理解我，因为理解很重要”等循环理由，或只有无关故事。')
h('如需计分',2)
p('六个词语各1分，共6分；匿名表格按四点各1分，共4分：时间限制与监督、相应感受、从爸爸角度思考、联系学习及亲子关系顾虑；开放回应按上表四项各1分，共4分。合计14分，由教师核对，不是官方分值。各意义点独立计分，漏一点不清零其他点。')
p('网页首次正确率只计客观点选、取证与词语任务，不是14分总评，也不代表口头概括能力。匿名表格与开放题不自动判对错；自检勾选只表示学生自查。')
h('教参中的口头追问',2)
p('你对辅导员的建议有什么看法？除了辅导员的建议，你认为还有哪些方法可以帮助这三位同学？每次可选一位同学，必须说为什么。合理建议包括：选合适时机问候、提前告知行程、商量学习与娱乐安排；须对应具体烦恼。')
new('课堂操作与记录保存')
for title,text in [
 ('逐段阅读','学生打开reading.html，左侧原文，右侧对应漫画与词语积累。六条消息按原课本顺序推进，发言者和时间保留。原漫画直接展示，无解锁步骤。'),
 ('原句取证','点原句，再点“加入原句”，连续加入会累加。也可以复制、粘贴和编辑。国明建议要兼顾直接表达与具体方法；美华建议要兼顾体谅与沟通。'),
 ('反馈与订正','首次错误只提示方向，5秒后可重试，再次错误指向关键词。首次答案不会被订正覆盖。主观表格与开放回应先提交再比较参考，教师核对。'),
 ('个人词语本','32项语境词库供按需查询，不强制收藏。学生自输词也能建卡，未收录词标待补释义。站内同浏览器的个人收藏按年级、单元、课文汇总，保留已有笔记和熟悉状态。'),
 ('学习记录','检查页填姓名、班级、学号或组别。可下载完整JSON、逐题CSV或文字记录，含首次回答、每次尝试、时间、自检及订正。系统不会自动发送给教师，也不上传Google Sheet。'),
 ('设备与备份','记录只在当前设备、当前浏览器。清理浏览器前导出；共享设备可分别使用个人浏览器资料或各自备份。离线文件与线上页面不自动互通。重置本课不删除个人词语本。'),
 ('资源与投影','28页课件原生文字至少30pt，题目页与讲解页分开，教师提示放备注。线上投影支持前后页、菜单和全屏。整套下载含可编辑PPT、教师DOCX、独立学生HTML、原漫画及离线投影；解压后保留slides目录。'),
 ('教师入口','教师材料为公开页面，链接不构成身份保护。课堂发送学生直达链接。学生若有不愿分享的家庭困扰，可改用假设情境；需要真实帮助时鼓励私下找教师或学校辅导员。')]:
 h(title,2);p(text)
doc.save(D/'downloads/teacher-guide.docx')
print('Saved teacher guide')
