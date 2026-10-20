from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[2];D=ROOT/'du/textbook/2hcl/unit2/zhu-yang-ye-sheng-dong-wu';HERE=Path(__file__).parent
L=json.loads((HERE/'lesson.json').read_text());M=json.loads((HERE/'teaching_meta.json').read_text());S=json.loads((D/'slides.json').read_text())
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for b in doc.styles.element.xpath('.//w:pBdr'):b.getparent().remove(b)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Noto Serif CJK SC';st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=Pt(16);doc.styles['Normal'].paragraph_format.space_after=Pt(5)
for name,size in [('Title',25),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:doc.styles[name].font.size=Pt(size)
for name in ['Heading 1','Heading 2','Heading 3']:doc.styles[name].paragraph_format.space_before=Pt(9);doc.styles[name].paragraph_format.space_after=Pt(5)
footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('华文通　助养野生动物　教师版　');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(9)
def p(s,bold=False,kai=False):
 a=doc.add_paragraph();r=a.add_run(str(s));r.bold=bold
 if kai:r.font.name='LXGW WenKai TC';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'LXGW WenKai TC')
 return a
def h(s,level=1):doc.add_heading(s,level)
def new(s):doc.add_page_break();h(s)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,v in enumerate(headers):t.rows[0].cells[i].text=str(v)
 for vals in rows:
  cells=t.add_row().cells
  for i,v in enumerate(vals):cells[i].text=str(v)
 head=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(head)
 for ri,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','bottom','left','right']:
    b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
   pr.append(borders);margin=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margin.append(e)
   pr.append(margin)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE8EE');pr.append(sh)
   for a in c.paragraphs:
    a.paragraph_format.space_after=Pt(3);a.paragraph_format.line_spacing=1.15
    for r in a.runs:r.font.size=Pt(10.5);r.bold=ri==0
 return t
doc.add_paragraph('助养野生动物 教学流程与参考答案','Title')
p('中二高级华文　单元二　生活空间　1课时',True)
p('本课训练信息定位、筛选概括和有理由的开放回应。教材原文是模拟网页，以助养计划和动物介绍组成；按阅读顺序划为七个教学段，逐段读文、取证。以下答案根据课文拟定，属于课堂参考。')
h('核心目标',2)
for i,o in enumerate(L['objectives'],1):p(f'{i}. {o}')
h('50分钟课堂流程',2)
table(['时间（分钟）','环节','操作'],[[r[0].replace('分钟',''),r[1],r[2]] for r in M['flow']],[1.0,1.25,4.85])
p('8组任务，32页可编辑PPT。第1—2页导入，第3—12页助养信息，第13—27页动物特点，第28—31页公益分享与回应，第32页检查。原生文字至少30pt；问题先出现，讲评下一页揭示。')
p('动物特点表按教参2.2以口头概括为主，网页保存原句依据，表格可按需要补写概括。正式打字集中在公益广告来源与3—4句回应。不加新短文、课后一周小测或实际募款。')
new('文本结构与图文对应')
table(['教学段','内容与主要信息','原漫画格号'],[[i+1,L['paragraphIdeas'][i],'、'.join(map(str,L['panels'][i]))] for i in range(7)],[.65,5.0,1.45])
p(M['boundary']);p(M['comic'])
h('学生困难与支架',2)
table(['困难','本课处理'],[['混淆对象','小标题定位，分清捐款用于动物，与助养者得到优惠。'],['只抄细节不会概括','先看身体、动作、性情；教师示范阿明，学生再说企鹅、可可与小哈。'],['忽略转折','先讲“虽然……但……”与“却”，再概括可可与小哈。'],['理由空泛','用建议、理由、例子、回扣组织3—4句，例子要能支持建议。']],[1.5,5.6])
p('O水准能力对应：按题干限定找信息、多点作答、语境词义、筛选与概括、开放评价。本课不硬加人物形象、线索或说明方法题。')
def task(t,i):
 h(str(i)+' '+t['title'],2)
 p(t['source']+'　课堂身份：'+t['role'])
 for q in t['questions']:p(q,True,True)
 p('参考：'+t['answer'],kai=True)
 p('依据：'+('；'.join(t['evidenceQuotes']) if t['evidenceQuotes'] else '学生所选真实广告；示例使用NParks公益海报。'),kai=True)
 p('答案点：'+'；'.join(t['points'])+'。')
 p('思考：'+t['explain'])
 p('易错：'+t['reject'])
 p('提示1：'+t['hints'][0]+'　提示2：'+t['hints'][1])
for idxs in [[0,1],[2,3],[4,5],[6]]:
 new('随段问题与讲评 '+str(idxs[0]+1)+(' 至 '+str(idxs[-1]+1) if len(idxs)>1 else ''))
 for i in idxs:task(L['tasks'][i],i+1)
 if idxs==[6]:
  h('五个课本词语题核对',2)
  table(['题号','释义','课文词语'],[['1','提供救助，帮助脱离危险。','拯救'],['2','迅速而灵敏。','敏捷'],['3','来回活动。','穿梭'],['4','看见好吃的食物就想吃。','馋'],['5','一种体育运动，双方相抱，运用力气和技巧，谁能使对方倒下谁就能获胜。','摔跤']],[.7,5.3,1.1])
new('动物特点表与评价')
p('课本原题：阅读课文，填写以下表格。列名为“动物”“概括主要特点”。按教参建议先口头概括，再核对。以下四行供教师讲评。',True)
table(['动物','概括主要特点'],[[n,L['tasks'][i]['answer'].replace('词语：敏捷、穿梭。','').replace('词语：馋。','').replace('词语：摔跤。','')] for i,n in [(3,'红毛猩猩'),(4,'非洲企鹅'),(5,'亚达伯拉象龟'),(6,'马来熊')]],[1.2,5.9])
h('课堂评分建议',2)
p('教材截图没有提供分值。需要计分时：词语题每词1分，共5分；动物表按独立意义点各1分，红毛猩猩3分，企鹅3分，象龟4分，马来熊4分，共14分。课堂不强制总分，网页只显示取证与点选首次正确率。')
p('特点接受等义表达，如“喜爱新鲜事物”对应探索，“喜欢甜食”对应最爱甜的东西。红毛猩猩的亲善大使身份可补充，不代替身体、动作、性情。象龟若只答“很重”可算重量这一意义点；缺少120公斤时教师可追问单位与数字。')
p('各意义点独立判断，不因漏一项而清零其他内容。原句可用于证明定位，口头或书面概括应删重复、保留关键关系。自动检查原句包含相关词句，不能证明学生已经会概括。')
new('公益广告分享与开放回应')
t=L['tasks'][-1]
for q in t['questions']:p(q,True,True)
p('第一句为课本原题。行动建议、为什么及3—4句书面回应为教师补充，用来落实本配套的必做开放回应。')
p('12分钟安排：浏览或挑选广告3分钟，同桌各分享约1分钟共2分钟，独立写回应5分钟，按要点自检2分钟。教师可课前打开官方资源页，减少全班搜索等待。')
table(['步骤','学生应写什么','教师判断'],[['建议','愿意实践哪项行动','对象清楚，行动可行。'],['理由','为什么这项行动有帮助','解释具体需要或影响。'],['例子','广告内容或具体情境','材料真实，支持理由。'],['回扣','这怎样帮助动物','解释联系，不只重复口号。']],[.7,2.5,3.9])
p(t['answer'],kai=True)
p('参考评价：每项0—1分，共4分，由教师按意义判断。广告名称、网址或发布者用于核对来源，不以拼写或格式替代理解。不同建议均可得分，不要求实际捐款。具体情境可以是假设，但不能冒充亲身经历。')
p('首次反馈：有没有说明为什么？再次提示：广告中哪一点能支持建议？最后对照：看例子与建议之间是否有解释。网页只保留回答并提供自检，不自动判开放题正确。')
h('已核对的公益宣传资源',2)
p('发布者：新加坡国家公园局 NParks。官方资源页：'+M['adurl'])
p('进入“Posters”下载3张海报，第3页为“Wildlife and Us”。其中“不喂食野生动物”及“人的食物可能使动物生病”的提示，是示例依据。资源页于2026年10月20日核对。也可用学生自行查到的真实公益广告。')
new('课堂操作与记录保存')
for s in ['学生入口为reading.html。读文页左侧原文、右侧对应漫画及词语积累，一次一段；读完可点“本段问题”。','问题左边显示对应原文。点原句，再点“加入原句”，第二句会累加；也可复制、粘贴和修改。首次答错给提示，5秒后可重试。','四种动物先口头概括，再选原句，最后对照。不能把取证正确率当成概括能力得分。开放回应必做，教师核对。','在“检查”页填姓名、班级、学号或组别，可下载完整JSON备份、逐题CSV或文字记录。首次作答、订正和尝试时间保留。','学习记录仅保存在当前设备与浏览器，不自动上传Google Sheet，不自动发给教师。清理浏览器前先导出；共享设备应分别使用个人浏览器资料或各自备份。','个人词语本可跨站内课文汇总，按年级、单元、课文筛选。同词保留原有笔记与熟悉状态。未知词仍可收藏，释义待补充。离线页与在线页不自动互通。','整套下载包含PPT、教师方案、独立学生HTML、原漫画、离线投影。外部公益海报需联网，建议课前打开或自行下载备用。','教师资料为公开页面。给学生发送学习页直达链接；教师入口未设置身份验证。']:
 p(s)
h('本课实施范围',2)
p('本课复用华文通通用引擎与共享词语本。中二单元五、六旧课文按既定安排暂缓处理。课件、教师参考、在线投影与下载包均以同一课文内容生成。')
doc.save(D/'downloads/teacher-guide.docx')
print('Saved teacher-guide.docx')
