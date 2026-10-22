from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1];O=R/'repo/du/textbook/1hcl/unit3/xingyao-changi';L=json.loads((O/'lesson.json').read_text())
D=Document();sec=D.sections[0];sec.page_height=Inches(11);sec.page_width=Inches(8.5);sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=D.styles[name];st.font.name='Noto Serif CJK SC';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.font.color.rgb=RGBColor(0,0,0)

for sty in D.styles:
 for tag in list(sty._element.xpath('./w:pPr/w:pBdr')):tag.getparent().remove(tag)
D.styles['Subtitle'].font.italic=False
D.styles['Normal'].font.size=Pt(11);D.styles['Normal'].paragraph_format.space_after=Pt(4);D.styles['Normal'].paragraph_format.line_spacing=Pt(16)
for name,size in [('Title',23),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:D.styles[name].font.size=Pt(size)
def p(t,style=None,kai=False):
 z=D.add_paragraph(t,style);z.paragraph_format.widow_control=True
 if kai:
  for r in z.runs:r.font.name='Huawentong WenKai';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Huawentong WenKai')
 return z
def h(t,new=False):
 z=D.add_heading(t,1);z.paragraph_format.page_break_before=new;z.paragraph_format.space_before=Pt(8);z.paragraph_format.space_after=Pt(5);return z
def table(headers,rows,widths):
 t=D.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 tag=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(tag)
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 for ri,row in enumerate(t.rows):
  no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
  for i,c in enumerate(row.cells):
   c.width=Inches(widths[i]);pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   v=OxmlElement('w:vAlign');v.set(qn('w:val'),'center');pr.append(v)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DDEBF5');pr.append(sh)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=Pt(15)
    for r in par.runs:r.font.size=Pt(11);r.bold=ri==0
 return t
p('魅力四射 星耀樟宜','Title');p('教师教学流程与参考答案','Subtitle');p('中一高级华文  单元三  核心课文  2—3课时')
p('本课训练学生通过数字和例子理解事物特点。按7个自然段依次读懂，随段完成课本五行统整表；最后独立解释数字的作用，并向朋友推荐三个地方。建议三课各40分钟。')
h('教学目标')
for i,x in enumerate(L['objectives'],1):p(str(i)+'. '+x)
h('三课时课堂流程')
rows=[['第一课\n40分钟','第1—3段','导入与第1段8分钟；第2段14分钟；第3段15分钟；检查3分钟。示范小标题与第一行表，再多句取证。'],['第二课\n40分钟','第4—7段','森林谷10分钟；设施与短答15分钟；美食8分钟；总结5分钟；订正2分钟。独立写25米的作用。'],['第三课\n40分钟','全文整理与推荐','词语12分钟；推荐21分钟（构思4、独立写10、对照交流7）；检查、词语整理与导出7分钟。']]
table(['课次','内容','教学安排'],rows,[.8,1.15,5.15])
h('两课时安排')
p(L['twoLessonFlow'])
h('文本结构与图文对应',True)
table(['段落','内容与作用','原漫画格号'],[(str(i+1),L['paragraphIdeas'][i],str(L['panels'][i][0])+'—'+str(L['panels'][i][-1])) for i in range(7)],[.55,5.45,1.1])
h('难点与支架')
p('数字多，容易混淆对象：数字、单位、对象一起读，特别区分长度与高度、数量与品种、商店总数与餐馆数量。')
p('找到数字却不会分析作用：先示范第2段，再共同分析雨漩涡，最后独立写步行网的作用。方法名不能代替具体特点。')
p('推荐只有“很好玩”：用建议、理由、文本依据与回扣，说明特点怎样适合朋友的需要。允许不同推荐，不虚构个人经历。')
h('教材与图片边界')
for x in L['boundaries']:p(x)
Q={
'intro':'第1段：“不过，如果你以为它只是一个普通的购物中心，那就太小看它啦！”',
'exterior':'第2段：“它的外壳由大约1万8000根钢铁制成的横梁构成，覆盖着9000多片特制的玻璃。”；“都是按照它所处的位置量身定制的。”',
'waterfall':'第3段：“它高40米，从五楼一直到地下二楼”；“这些水通过水管可循环使用，因此不会造成浪费。”；“还能帮助降低室内温度。”',
'forest':'第4段：超过2000棵树木、超过10万株灌木、约120个品种；来自西班牙的两棵超过100岁的橄榄树；访客可以松懈身心。',
'play':'第5段：“好玩又刺激”；天悬桥的云雾、镜子迷宫的奇幻神秘、天空之网的步步惊心；步行网长50米、离地25米，蹦跳网长250米。',
'height-effect':'第5段：“步行网是一条50米长的网状步道，距离地面25米高，走起来摇摇晃晃的，需要几分胆量”。',
'food':'第6段：“这里有200多家商店，其中不少和‘吃’有关，零食店、甜品店、饮料店和特色餐馆应有尽有。”',
'ending':'第2—6段五方面特点；第7段“不仅为国人和旅客提供了一个放松身心的好去处，而且也让樟宜机场这一世界级的航空中心变得更加迷人，更有魅力”。',
'word-families':'阅读放大镜一原题及木、桥、顶、习四组例子。属于字义与构词，不强套课文证据。',
'recommend':'课本读后说；依推荐对象从第2—6段选择准确数字或例子。答案不唯一。'}
for i,t in enumerate(L['tasks']):
 h(t['title'].replace(' · ',' '),i==0)
 p(t['role']+'。'+t['source'])
 for q in t['questions']:p(q,kai=True)
 p('目标：'+t['target'])
 p('文本依据：'+Q[t['id']],kai=True)
 p('参考答案：'+t['answer'],kai=True)
 p('信息加工与接受范围：'+t['explain'])
 p('常见失分：'+t['reject'])
 p('参考评价：'+t['scoreGuide'])
 p('提示一：'+t['hints'][0]+'  提示二：'+t['hints'][1])
 if t['checks']:
  p('选择题诊断：'+'；'.join(c['question']+' 答'+chr(65+c['correct'])+'。'+' '.join(chr(65+j)+'项：'+x for j,x in enumerate(c['why'])) for c in t['checks']))
 for q,a in t.get('oral',[]):p('口头追问：'+q+'  参考：'+a)
h('五方面统整表参考')
rows=[]
for t in L['tasks']:
 if 'answerTable' in t:rows.append([t['answerTable']['name']]+[f['correct'] for f in t['answerTable']['fields']])
table(['说明对象','事物的特点','对特点的说明','说明的方法'],rows,[1.15,1.55,3.1,1.3])
p('补充第4段的完整信息：灌木超过10万株，两棵橄榄树超过100岁。第5段还须读出蹦跳网长250米和腾空而起的感觉。网页表格是概括，完整原文始终可回看。')
h('词语串串串的课本示例')
table(['共同字','课本例子','理解方向'],[['木','树木、灌木、草木、木材、木柴','与树木有关'],['桥','天悬桥、波浪桥、木桥、铁桥、吊桥','一类跨越通道'],['顶','屋顶、房顶、山顶、楼顶、头顶','上端或最高部分'],['习','学习、练习、补习、预习、复习','学习或练习']],[.6,4.55,1.95])
h('学生作答的教师覆核')
p('25米作用题按3个独立意义点检查；推荐题按每处3点、共9点作教师参考。词语题允许其他正确词语。不把不同题型强行合成一个自动总分。')
h('操作与数据记录',True)
p('学生从课文主页进入逐段图文阅读。当前段原文在左、原漫画与词语收藏在右；任务页保持左文右题。证据框可累加两句及以上原文，并支持复制、粘贴与编辑。')
p('答错后先显示提示，5秒后可重试；再次错误后可以看讲解。示范在学生提交后出现。首次回答、尝试与订正分别保留。主观题只检查是否填写，不把结构完整当作理解正确。')
p('检查页显示7组客观任务的首次全部答对率。客观部分有8个四选一判断、15个表格点选、1个多句证据框。3组主观任务为作用题、词语串串串和推荐，均由教师核对。师生共构任务的完成不等于学生独立掌握。')
p('页面上方可填姓名、班级、学号。普通入口在同设备、同浏览器保存；检查页可导出完整JSON备份、逐题CSV与文字记录。重置本课不删除跨课个人词语本。清浏览器资料前先导出。')
p('通过现有课文课堂看板选择本课、创建班级任务并发放学生专属链接，可沿用平台收件。普通链接与离线包不自动发给教师。教师页与下载材料公开包含答案，不构成身份保护。')
p('词语本有44条经语境核对的预设释义，只有学生收藏的词进入个人词语本。未知词也能建卡，但标明释义待补充。跨课汇总限同域、同浏览器；离线与在线记录分开。')
p('可编辑课件与在线投影页序一致。原漫画按五页完整保留，课件仅通过图片框显示某一格。整套包附字体；缺字体的设备也可直接使用在线投影。离线学生页可直接打开reading.html；投影若从本地使用，需要HTTP服务。')
h('教材题覆盖与教学顺序')
p('预习二的独一无二、腾云驾雾、与时俱进，分别随第2、5、7段处理。阅读放大镜二五行表随第2—6段完成，读完第7段才汇总。阅读放大镜一六项构词与读后说三个推荐，在全文阅读后完成。')
p('没有新增短文、课后一周小测、速度奖励或个人排名。视频导入可由教师在原SLS资源中按需使用；本配套以已有漫画和生活经验完成导入。')
h('资料来源')
p('教师上传的课文截图：标题页、第68页起正文、阅读放大镜一与二、读后说；教师手册两页教学流程；五页《魅力四射 星耀樟宜》原漫画。参考答案与评分为依据原文和教参编制的教学参考，不冒称官方考试评分。')
D.save(O/'downloads/teacher-guide.docx');print('DOCX saved')
