"""Editable teacher flow and answer guide, with fixed pedagogical page breaks."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[2];D=R/'du/textbook/1hcl/unit3/guojia-meishuguan'
L=json.loads((D/'lesson.json').read_text());doc=Document()
sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.72)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Noto Serif CJK SC';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal'];normal.font.size=Pt(11.5);normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(7)
for nm,sz in [('Title',24),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
 st=doc.styles[nm];st.font.size=Pt(sz);st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(6)
def para(s,bold=False):
 p=doc.add_paragraph();r=p.add_run(s);r.bold=bold;return p
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for r in rows:
  cells=t.add_row().cells
  for i,s in enumerate(r):cells[i].text=str(s)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
  if ri==0:
   rr=OxmlElement('w:tblHeader');trpr.append(rr)
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);pr=c._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders);m=OxmlElement('w:tcMar')
   for side,val in [('top','80'),('bottom','80'),('left','100'),('right','100')]:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');m.append(el)
   pr.append(m)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DDEAF3');pr.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
    for r in p.runs:r.font.size=Pt(10.5);r.bold=ri==0
 return t
def page(title):
 p=doc.add_heading(title,1);p.paragraph_format.page_break_before=True
def task(n):
 t=L['tasks'][n-1];doc.add_heading(t['title'],2);para(t['source']+'；'+t['role']+'。')
 for q in t['questions']:para(q,True)
 para('参考答案　'+t['answer'])
 para('证据与思考　'+t['explain'])
 para('课堂评分参考　'+t['scoreGuide'])
 para('常见失分　'+t['reject'])
 para('先提示　'+t['hints'][0]+'　再定位　'+t['hints'][1])
 if t.get('oral'):
  for q,a in t['oral']:para('口头追问　'+q+'　参考：'+a)

doc.add_heading('国家美术馆探秘 教师教学手册',0)
para('中一高级华文　第三单元　巩固课文　1—2课时',True)
para('本课围绕八段说明文，帮助学生把数字和例子连到事物的特点，解释五个“秘密”，并用准确的课文内容支持自己的选择。先读懂当前段落，再作答与订正，读完全篇后才进行统整。')
doc.add_heading('核心目标',1)
for s in ['准确概括建筑与画作中的五个“秘密”，分清原文明写与合理推测。','把数字、单位或例子与所说明的特点相连，解释方法的具体作用。','解释“既熟悉又陌生”的两面意思，保持同一对象进行仿写。','完成必做开放回应，用具体内容说明理由及其与立场的关系。']:para('• '+s)
doc.add_heading('两课时安排',1)
table(['课次','课堂流程','时间'],[
['第一课','导入／第1—2段／第3—6段／第7段书答与讲评／订正和词语整理','3＋10＋18＋14＋5＝50分钟'],
['第二课','回顾／第8段与开放回应／全文统整／仿写与查找／检查订正','4＋14＋12＋15＋5＝50分钟']],[.65,4.7,1.65])
para('第一课文字作答1项：第7段解释题。第二课文字作答2项：开放回应和仿句。其余主要采用点选、取证、表格配对和口头追问。')
doc.add_heading('一课时紧凑安排',1)
para('导入2分钟，图文自读10分钟，按段核对12分钟，统整6分钟，仿句与查找8分钟，必做开放回应8分钟，检查4分钟，共50分钟。重复的事实小题采用口头核对，第7段可先口述；保留开放题书写和全部课本练习。程度较弱班级宜分两课。')
para('不另加短文、课后一周小测、打卡或速度排名。课堂时间用于理解和订正。')

page('段落证据与图文对应')
table(['段号','内容与证据焦点','原漫画格号'],[[i+1,L['paragraphIdeas'][i],'、'.join(map(str,L['panels'][i]))] for i in range(8)],[.55,5.1,1.35])
doc.add_heading('全文结构',2)
para('第1—2段总介与引入探秘；第3—7段逐一介绍建筑和画作的五个秘密；第8段说明美术馆的价值并呼吁读者亲自探索。教学顺序与段落一致。')
doc.add_heading('学生困难与支架',2)
para('词义：殖民地、自治、原貌、氧化、里程碑等提供语境中英文释义。关系：数字要连单位，例子要连概括句。表达：先分清“看过图案”和“了解画作”，开放题再用立场、理由、例子与回扣。')
doc.add_heading('事实边界',2)
para('“6万4000平方米”“9000余件”等保留为课文数据。“自治”不改为“独立”；“部分原貌”不扩大为全部。第6段未写圆顶原本的具体颜色，漫画色变是示意，不能作为补写原色的证据。')
para('课文与原漫画的画家姓名为“钟泗滨”，本配套照录；南艺等官方资料也写作“钟泗宾”，英文同为Cheong Soo Pieng。画作示意不能代替《晒咸鱼》原画。')

page('第一段')
task(1);page('第二段');task(2)
page('第三段')
task(3);page('第四段');task(4)
page('第五段')
task(5);page('第六段');task(6)
page('第七段')
task(7);page('第八段与开放回应');task(8)
page('全文统整与课本练习')
doc.add_heading('9 五个秘密与课本统整表',2)
para('课本原题：国家美术馆藏了哪些秘密？阅读课文，了解课文内容，通过数字和例子理解事物的特点。',True)
para('五个秘密：两座历史建筑合建的来历；地下通道的用途；窗户变成门；圆顶因氧化而变色；钞票上的《晒咸鱼》。先口头概括，再完成表格。')
table(['对象','特点','具体说明','方法'],[['建筑','面积很大','6万4000平方米，约9个足球场','列数字，结合比较'],['建筑','历史意义重要','政府大厦见证日本投降和宣布自治','举例子'],['建筑','保留部分原貌','法官座位、被告栏、棕色天花板','举例子'],['画作','馆藏丰富','9000余件收藏品','列数字']],[.6,1.3,3.9,1.2])
para('课堂参考：12格各1点。数字中的“平方米”和“件”不能调换；每格独立判断，不因一格错误而全行清零。“9个足球场”可说明为列数字结合比较，不强迫学生只记一个标签。')
page('课本仿写与资料查找');task(10)

page('课堂操作与资料来源')
doc.add_heading('学生与教师入口',2)
para('学生打开“互动学习”，可从上次位置继续。读文时一次一段，原文与对应原漫画直接显示；红色粗体下划线词语可点开中英文释义，并加入跨课文词语本。词语可用闪卡复习。')
para('任务页左边原文、右边当前问题。证据框可连续加入多句，也可复制、粘贴和编辑。第一次答错先给提示，5秒后可订正；第二次可继续查提示或展开讲解。开放题与仿句提交后对照，自检不等于自动评分。')
para('教师可直接开PPT或在线投影。问题与关键答案分开显示。统一讲解时先让学生独立作答，再展示下一页；避免提前给出统整表和开放题示例。')
doc.add_heading('记录与下载',2)
para('普通链接使用当前浏览器本地保存，保留姓名、班级、学号、首次作答、尝试和订正。检查页可导出JSON完整备份、CSV逐题明细和文字记录。重置本课不删除个人词语本。清浏览器数据前先导出；离线页关闭前也要备份。')
para('如要在教学管理助手查看班级进度，先建立本课班级任务，再发放各学生专属链接。公共学习链接不会自动交给老师。本配套不新增班级或录入学生名单。')
para('整套包包括网页、四页原漫画与逐格显示副本、独立离线学习页、可编辑PPT、教师手册及投影图片。离线页面的查找资料链接需要网络；无网络时使用下列已核实的教师答案。')
page('查找题核对与资料来源')
para('50新元钞票背面的另一幅画来自陈文希（Chen Wen Hsi），题名为Two Gibbons Amidst Vines，画面是藤蔓间的两只长臂猿。MAS资料简称Gibbon。不要把钞票整体设计者与画作作者混淆。')
para('核对日期：2026年10月23日。以下为查找题资料，不替代附件课文。')
for s in [
'新加坡金融管理局，Circulation Currency: Notes，Portrait Series $50部分：\nhttps://www.mas.gov.sg/currency/circulation-currency/circulation-currency-notes',
'国家美术馆，FY2016 Annual Report，印刷页16（PDF第19页）：\nhttps://www.nationalgallery.sg/content/dam/about/annual-reports/reports/FY2016%20Annual%20Report.pdf',
'南洋艺术学院，Cheong Soo Pieng作者资料：\nhttps://iseaa.nafa.edu.sg/resources/collections/artist/cheong-soo-pieng']:para(s)
doc.core_properties.title='国家美术馆探秘 教师教学手册';doc.core_properties.author='华文通'
doc.save(D/'downloads/gallery-teacher.docx');print('Teacher DOCX created')
