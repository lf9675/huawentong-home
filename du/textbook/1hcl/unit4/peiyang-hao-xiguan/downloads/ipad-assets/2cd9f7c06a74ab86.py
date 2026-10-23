"""Editable teacher flow and answer guide, with fixed pedagogical page breaks."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[2];D=R/'du/textbook/1hcl/unit4/peiyang-hao-xiguan'
L=json.loads((D/'lesson.json').read_text());doc=Document()
sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.72)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Noto Serif CJK SC';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=Pt(17);normal.paragraph_format.space_after=Pt(6)
for nm,sz in [('Title',24),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
 st=doc.styles[nm];st.font.size=Pt(sz);st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(6)
def para(s,bold=False):
 p=doc.add_paragraph();r=p.add_run(s);r.bold=bold;return p
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for r in rows:
  cells=t.add_row().cells
  for i,s in enumerate(r):cells[i].text=str(s)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
  if ri==0:
   rr=OxmlElement('w:tblHeader');trpr.append(rr)
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);pr=c._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders);m=OxmlElement('w:tcMar')
   for side,val in [('top','80'),('bottom','80'),('left','100'),('right','100')]:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');m.append(el)
   pr.append(m)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DDEAF3');pr.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
    for r in p.runs:r.font.size=Pt(10.5);r.bold=ri==0
 return t
def page(title):
 p=doc.add_heading(title,1);p.paragraph_format.page_break_before=True
def task(n):
 t=L['tasks'][n-1];doc.add_heading(t['title'],2);para(t['source']+'；'+t['role']+'。')
 for q in t['questions']:para(q,True)
 para('参考答案　'+t['answer'])
 para('证据与思考　'+t['explain'])
 para('课堂评分参考　'+t['scoreGuide'])
 para('常见失分　'+t['reject'])
 para('先提示　'+t['hints'][0]+'　再定位　'+t['hints'][1])
 if t.get('oral'):
  for q,a in t['oral']:para('口头追问　'+q+'　参考：'+a)

brief=json.loads((D/'teaching-brief.json').read_text())
doc.add_heading('培养好习惯',0)
para('中一高级华文　单元四·生活空间\n教师教学流程与答案　一课时（50分钟）',True)
para('课文是一份青年节活动海报。学生先读八种好习惯，再比较互动剧场与工作坊，最后用真实经验说明怎样维护友情。按海报内容顺序推进，不人为编造段落编号。')
doc.add_heading('学习目标',1)
for x in brief['targets']:para('• '+x)
doc.add_heading('课堂安排',1)
table(['环节','分钟'],brief['flow'],[5.9,1.1])
para('书面作答只有两项：工作坊的目的、维护友情的经验。四个词义题采用口头共构；其余课本信息题先点选，再口头补成完整句子。')
para('课时按50分钟设计。若实际课时较短，压缩导入和全班分享，保留两项书答与订正。旧词语页按需复习，不要求整套再做一遍。')

page('原课本任务与图文顺序')
para('课本第一题：用自己的话解释“坚持不懈、开卷有益、一诺千金、一帆风顺”。教学按海报顺序穿插口头完成，保留四词，不另加书写题。')
para('课本第二题：互动剧场有什么特别之处？举办工作坊的目的是什么？工作坊适合哪些人参加？如何报名参加青年节特备活动？')
para('课本第三题：你是怎么维护友情的？和同学分享你的经验。“为什么有帮助”作为表达支架，不冒充课本原题。')
table(['海报内容','原漫画格号'],[[a,'、'.join(map(str,b))] for a,b in brief['mapping']],[4.4,2.6])
para('四页原漫画共32格均保留。第32格是全篇总结，保留在完整漫画重读处；主学习页按内容选用相关格子。图中的补充故事帮助想象，不作为海报事实的答题证据。')
para('资料差异：课本海报的报名网址为www.huanlecommunity.com.sg，教参与漫画第32格使用另一网址。答案遵照课本原图，原漫画不改动。日期、嘉宾与网址只作教材阅读资料。')

page('四个词义题：口头共构')
table(['词语','最低意义要求','常见误解'],[
['开卷有益','读书有好处。','“开卷”不是只打开书，也不要求必须读完一本才有益。'],
['坚持不懈','坚持努力，不松懈、不轻易放弃。','只说做得快或做得好，没有说明持续努力。'],
['一诺千金','说话算数，信守承诺。','“千金”不等于一定牵涉很多钱或只有重大承诺才适用。'],
['一帆风顺','事情进行得很顺利，没有阻碍。这里指朋友相处顺利。','把友谊的小船理解为真的一起坐船。']],[1.2,2.7,3.1])
doc.add_heading('口头提问顺序',2)
para('开卷有益：读“勤于阅读”，问读书有什么好处，再请学生用自己的话概括。')
para('坚持不懈：联系“滴水穿石”，问遇到困难时怎样做，突出持续、不放弃。')
para('一诺千金：联系“信守诺言”“言出必行”，问答应同学的事为什么要做到。')
para('一帆风顺：先理解顺风航行，再放回“友谊的小船”这一语境，解释朋友相处。')
doc.add_heading('学生支架',2)
para('允许先用简单中文说明，教师帮助整理成一句完整解释。必要时可查看词语弹窗中的英文释义，最后仍回到中文表达。四题不重复打字、不设速度排名。')
para('教师听意义是否完整，不以某一个关键词是否出现机械判分。能用不同措辞说出相同意思即可。')

page('互动剧场与工作坊目的')
task(4)
task(6)

page('参加对象与报名方式')
task(7)
task(8)
doc.add_heading('活动资料核对',2)
table(['活动','日期与时间','地点'],[['互动剧场','7月21日晚上7点至9点','欢乐民众俱乐部礼堂'],['工作坊','7月20日下午3点至5点','欢乐民众俱乐部礼堂']],[1.3,3.1,2.6])

page('维护友情：必做开放题')
task(9)
doc.add_heading('检查什么',2)
para('做法是否具体：如先倾听、兑现承诺、主动道歉、尊重不同想法。不能只有“对朋友好”这样笼统的口号。')
para('经验是否相关：交代一件真实发生的事，能看出自己做了什么。不要求经历特别精彩。')
para('解释是否连上友情：说明如何减少误会、建立信任或让朋友感到受尊重，理由须与自己的做法对应。')
doc.add_heading('反馈与订正',2)
para('教师可以反馈：“你已写出做法，请补充当时的具体情况”；或“经历很清楚，再说明这件事怎样帮助了你们的友情”。学生保留首次回答，再作订正。')
para('示例只是组织表达的参考，不要求所有学生写同一种做法，也不将示例当作学生真实经历。不得根据关键词自动判定开放题正确。')

page('课堂操作与旧资源整合')
doc.add_heading('互动学习',2)
para('原文与对应漫画按海报顺序显示。红色粗体下划线词语可点开中英文释义，并加入按年级、单元、课文汇总的个人词语本与闪卡。图片沿用原图，逐格显示只作无损裁切。')
para('任务页左侧原文、右侧问题。工作坊目的题可连续选取多句证据，再编辑证据框。客观题答错后先看提示，5秒后可订正；第二次后可展开讲解。书面题由教师覆核。')
doc.add_heading('旧词语页的处理',2)
para('保留原网址和16个词，将旧版38道重复较多的练习整合为16道语境用法题。修正过窄或绝对化的解释，并为每题提供理解提示与答案说明。题目属于巩固练习，不冒充课本原题。')
para('已经理解的词不必重新练。课堂先完成主课，在检查时收藏不熟悉的词，再按需要打开旧词语入口。两页共享同一课文的个人词语本，作答进度分别保存。')
doc.add_heading('记录与下载',2)
para('普通链接在当前设备同一浏览器本地保存姓名、班级、学号、首次作答与订正。检查页可导出完整JSON备份、逐题CSV和文字记录。公共链接不会自动把答案交给教师。')
para('可编辑PPT共25页，问题与参考分开。在线投影按同样顺序显示。教师可先停在问题页，让学生独立回答，再展示参考。')
para('下载配套包含离线学习页、网页、原漫画、投影图片、可编辑PPT、教师手册及制作源文件。离线使用结束前请导出记录。旧词语页也随整套包提供，网页中的平台首页链接需要网络。')
doc.add_heading('资料来源',2)
para('教师附件：课本海报两页、课本问题页、教学参考两页、四页原漫画。答案以课本海报为依据，教参辅助确定意义要求。制作日期：2026年10月23日。')
doc.core_properties.title='培养好习惯 教师教学流程与答案';doc.core_properties.author='华文通'
doc.save(D/'downloads/habits-teacher.docx');print('Teacher DOCX created')

