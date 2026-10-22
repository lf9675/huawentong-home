from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent.parent;L=json.loads((R/'lesson.json').read_text())
d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=Inches(.68);s.left_margin=s.right_margin=Inches(.72)
for name in ['Normal','Body Text','Title','Heading 1','Heading 2','Heading 3']:
 st=d.styles[name];st.font.name='Noto Serif CJK SC';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.font.size=Pt(11.5);st.paragraph_format.space_after=Pt(4);st.paragraph_format.line_spacing=Pt(16)
for name,size in [('Title',26),('Heading 1',19),('Heading 2',15),('Heading 3',12.5)]:d.styles[name].font.size=Pt(size);d.styles[name].font.color.rgb=RGBColor.from_string('185E94')
for name in ['Title','Heading 1','Heading 2','Heading 3']:
 d.styles[name].paragraph_format.line_spacing=1.1;d.styles[name].paragraph_format.space_before=Pt(8);d.styles[name].paragraph_format.space_after=Pt(4)
def p(text,style=None):return d.add_paragraph(text,style)
def head(text,level=1):d.add_heading(text,level)
def table(headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
 for c,text in zip(t.rows[0].cells,headers):c.text=text
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for row in rows:
  for c,text in zip(t.add_row().cells,row):c.text=str(text)
 if len(headers)==2:
  for i,c in enumerate(t.columns):c.width=Inches([1.4,5.5][i])
 for row in t.rows:
  if len(headers)==2:
   for i,c in enumerate(row.cells):c.width=Inches([1.4,5.5][i])
  for c in row.cells:
   for paragraph in c.paragraphs:paragraph.paragraph_format.space_after=Pt(2);paragraph.paragraph_format.line_spacing=Pt(16)
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   m=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');m.append(e)
   c._tc.get_or_add_tcPr().append(m)
 return t
def new(title):d.add_page_break();head(title)
def task(t,short=False):
 head(t['title'],2);p(t['role']+'；'+t['source'])
 for question in t['questions']:p(question)
 head('参考答案',3);p(t['answer'])
 if not short:head('证据与加工',3);p(t['evidenceText'])
 p('解释：'+t['explain']);p('常见失分：'+t['reject']);p('提示：'+'；'.join(t['hints']));p('评价：'+t['scoreGuide'])
p('世界各地的新年风俗','Title');p('中一高级华文 · 单元二 · 生活空间\n一课时（50分钟）教师流程与参考答案')
head('核心目标');[p(g) for g in L['goals']]
head('课堂流程');t=table(['环节','分钟','安排'],L['flow'])
for i,c in enumerate(t.columns):c.width=Inches([1.15,.5,5.25][i])
for row in t.rows:
 for i,c in enumerate(row.cells):c.width=Inches([1.15,.5,5.25][i])
p('共50分钟。8个核心任务，2处书面作答。进度较慢时，减少全班汇报人数，保留课本三题、独立短答和同桌分享。')
new('文本分析与教学安排')
p('这是一篇按地区并列介绍风俗的说明性图文材料。核心关系是“时间—迎新方式—寓意”。不加入人物分析、情节线索与修辞辨认。')
table(['课文位置','信息重点与词语障碍'],[(str(i+1)+' · '+L['paragraphLabels'][i],r['action']+'；'+r['meaning']+'。\n词句：'+r['obstacle']) for i,r in enumerate(L['rows'])])
head('学生难点与支架',2);p('学生容易把动作当寓意，把各地信息混淆。先按原文顺序逐段整理，读完整篇再概括共同愿望。用“什么时候／做什么／表示什么”帮助筛选，保留原文时间和条件。')
head('原题保留与拆步',2);p('第一题五组搭配完整保留。第二题原表右侧“象征意义”为共同栏；前五步辨认各地具体寓意，第六步再用两例归纳共同愿望。第三题保留查资料和同学分享，补问“为什么有趣、怎样表达愿望”。')
p('教材无段号，网页按地区原顺序编为1—5段。三页漫画保持原图原台词：段1对应1—3格，段2对应4—6格，段3对应7—9格，段4对应10—12格，段5对应13—16格。17—18格为元旦与春节小知识，整页回读保留。')
new('逐段参考：西班牙');task(L['tasks'][0]);new('逐段参考：古巴');task(L['tasks'][1])
new('逐段参考：意大利');task(L['tasks'][2]);new('逐段参考：丹麦');task(L['tasks'][3])
new('逐段参考：世界各地的华人');task(L['tasks'][4]);head('可接受补充',2);p('“不管工作多忙碌，路途多遥远，大家都要回家团聚”可说明珍惜亲情与团圆；这不是所有地区共同的做法。')
head('课本小知识',2);p('元旦：阳历新年的第一天，即1月1日。\n春节：农历新年的第一天，即正月初一。农历新年也叫春节。\n只核对名称与日期，不增加阴阳历或节气测试。')
new('综合短答：不同做法，共同愿望');task(L['tasks'][5]);head('意义得分与接受范围',2);p('“希望新年幸福顺利”“祝愿来年生活美好”等等义表达可接受。两个准确例子须来自两个地方；只列动作可获得内容方面的部分评价，还须解释愿望。各意义单位独立给分。')
head('证据边界',2);p('“象征、以示、表示、预示、希望”表达文化寓意与期盼，不能证明好运必定发生。漫画帮助理解，不能用原文没有的画面细节作为文字证据。')
new('词语搭配与检索分享');task(L['tasks'][6],True);head('第三题安排',2);p(L['tasks'][7]['questions'][0]);p('补问：'+L['tasks'][7]['questions'][1]);p('核对网站发布者，找时间、做法与寓意，记录出处，写三四句后与同桌分享。使用已核对资料入口，也允许自行检索可靠资料。\n看法—理由—例子—回扣：不能只写“因为很好玩”。')
new('开放题参考与资料核对');task(L['tasks'][7]);head('检索资料',2);p('日本国家旅游局（JNTO）：New Year food traditions。2018年12月21日发布，2026年10月22日核对。\nhttps://www.japan.travel/en/sg/story/new-year-food-traditions/')
p('示范只依据“新年前夜吃荞麦面，以细长形状象征长寿”，不加入有争议的起源年代。断网时教师可介绍此已核对资料与出处，学生练习解释；不因此另加课外作业。')
new('记录、评价与课堂操作')
head('评价口径',2);p('五个地区每组2个客观检查点，词语搭配5点，共15点。网页首次正确率按6个客观任务“整组全对”计算，与逐空得分不同。西班牙为教师示范，古巴为共同完成，不能把总正确率当成独立掌握率。两道短答各参考4个意义点，由教师覆核，不与自动统计混算。')
head('反馈与保存',2);p('第一次答错只给方向提示，5秒后重试；第二次进一步定位，可展开讲解。保留首次回答、尝试、时间、订正及自检。主观题提交后提供示范与自检，不自动打分。\n填写姓名、班级、学号后，可在“检查”页导出完整JSON、逐题CSV和文字记录。数据保存在同设备同浏览器。清除浏览器前须导出。尚无教师自动收件或Google Sheet上传。离线与在线作答分开。')
head('词语本与资源',2);p('27个词条可点选查义并收藏。自输未收录词时标“释义待补充”。词语本沿用共享格式，按年级、单元、课文筛选；重置本课不删词语本。\nPPT原生文字可编辑，在线投影同页序；可安装所附字体用于本机编辑。离线学习包含课文、原漫画、互动与词语本。教师页公开含答案。')
head('边界与审查',2);p('传统风俗按“课文介绍”理解，不说成当代每个人的实践。课堂不仿做高处抛物、泼水或触碰碎碗碟。三张原漫画不改内容。')
table(['关口','结果与依据'],[['课程与文本','通过：一课时，五段完整原文、三道原题保留。'],['语言与负担','通过：先逐段后综合，两处书面任务。'],['题目与答案','通过：区分做法与寓意，按意义评价短答。'],['互动与数据','功能验收检查保存、订正、多句取证、词语本和导出。'],['资源与维护','以lesson.json统一内容，复用共享引擎。']])
for el in d.styles.element.findall('.//'+qn('w:pBdr')):el.getparent().remove(el)
for r in d.element.findall('.//'+qn('w:r')):
 pr=r.get_or_add_rPr();fonts=pr.find(qn('w:rFonts'))
 if fonts is None:fonts=OxmlElement('w:rFonts');pr.insert(0,fonts)
 for k in list(fonts.attrib):
  if 'Theme' in k:del fonts.attrib[k]
 for k in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+k),'Noto Serif CJK SC')
d.core_properties.title=L['title']+' 教师流程与参考答案';d.core_properties.author='华文通';d.save(R/'downloads/teacher-guide.docx');print('DOCX written')
