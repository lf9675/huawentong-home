from pathlib import Path
import json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]
L=json.loads((R/'lesson.json').read_text())
d=Document();sec=d.sections[0]
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
for name,size in [('Normal',12),('Title',24),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
 s=d.styles[name];s.font.name='Noto Serif CJK SC';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
 s.paragraph_format.space_before=Pt(0)
 s.paragraph_format.space_after=Pt(5)
 s.paragraph_format.line_spacing=Pt(18 if size<=14 else 26)
 s.paragraph_format.keep_with_next=name.startswith("Heading")
 for node in s.element.xpath(".//w:pBdr"):
  node.getparent().remove(node)
def p(text,style=None):return d.add_paragraph(text,style)
def h(text):d.add_heading(text,1)
def note(label,text):
 x=p('');x.add_run(label+'　').bold=True;x.add_run(text);return x
def page():pass
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('华文通 · 中二高级华文 · ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
d.add_heading('多元而开放的新加坡',0)
p('中二高级华文｜单元六｜核心课文｜教师流程与参考答案')
h('教学目标与材料')
p('学生能概括七段内容及全文结构，解释事例与观点的关系，具体说明第五段的过渡作用，并用合理的具体例子说明多元与开放。')
p('依据本轮教师上传的课文、课本问题、教学参考及五页原漫画。教材的历史说法按课文理解，漫画属于情境示意，不能当作实物史料。原漫画五页40格完整保留；原文为作答依据。')
h('三课时流程（每课50分钟）')
for label,text in [('第一课','导入和第1段8分钟；第2段14分钟；第3段14分钟；第4段10分钟；口头回顾4分钟。'),('第二课','第5段12分钟；第6段20分钟；第7段8分钟；全文结构10分钟。'),('第三课','课本词语5分钟；四字短语12分钟；语境用词8分钟；开放回应18分钟；订正与词语本7分钟。')]:note(label,text)
h('两课时压缩（每课50分钟）')
p('第一课按上述第一课流程完成第1—4段。第二课：第5—7段及全文结构22分钟，课本词语3分钟，四字短语6分钟，语境用词4分钟，开放题12分钟，订正3分钟。压缩口头汇报；保留原题与开放回应。学生需要更多支架时采用三课时。')
h('课堂组织')
p('按“本段原文与漫画→本段问题→下一段”推进。先让学生尝试，再给提示和反馈。第2b题由教师完整示范取证与解释；第5题先解释过渡段，再让学生独立组织答案。全文综合任务放在七段学完之后。')
page();h('资源使用与记录')
for label,text in [
 ('学生入口','index.html 汇集逐段互动页、全文漫画回读、个人词语本和学习记录。互动页可连续加入多句原文，证据框可编辑。'),
 ('教师入口','teacher.html 汇集在线投影、61页可编辑PPT、教师DOCX、离线包和整套教学包。投影与PPT的题目、答案分开呈现，教师按进度翻页。'),
 ('作答反馈','客观题记录首次作答和订正；首次答错先提示，5秒后可再答。不设倒计时、速度分或排行榜。课堂正确率按整项任务核对，不作为考试分数。'),
 ('主观题覆核','完整填写是提交要求，不代表答对。教师逐一核对意义点、文本证据和解释，接受合理等义表达。开放题必须说明为什么，不能只列地点或食物。'),
 ('保存范围','记录保存在同域、同设备、同浏览器；刷新后可继续。可导出逐题CSV及完整JSON。没有自动上传教师或Google Sheet，也没有账号云同步。教师可按班级现有方式收取导出文件。'),
 ('词语本','词语按中二／单元六／本课归类，显示出处和释义。学生自行选择需要复习的词，可设置熟悉程度、用闪卡复习；不把全部词表强制收藏。'),
 ('离线使用','解压离线包后打开index.html。离线包可直接读文、做题和管理词语；在线与离线记录分开；词语本可导出导入，作答记录可导出留存。若浏览器限制本地文件保存，请先导出记录。'),
 ('旧资源处理','旧综合学习单和随堂题库的重复观点题、钱币题、排序题与词语题已合并。旧学生链接指向统一入口引导；旧源文件归档，既有后台数据保留。新题采用新版本，不迁移旧分数。')]:note(label,text)
h('图文对应')
p('第1段：第1—2格；第2段：第3—8格；第3段：第9—18格；第4段：第19—24格；第5段：第25格；第6段：第26—38格；第7段：第39—40格。')
groups=[['p1-view','p2-position'],['p2-coins','p3-timeline'],['p3-relation','p4-open'],['p5-link'],['p6-food','p6-art'],['p7-ending','whole-structure'],['words-prep','phrases'],['words-usage'],['open-response']]
for ids in groups:
 page()
 for id in ids:
  t=next(x for x in L['tasks'] if x['id']==id)
  h(t['title']);p(t['source']+'｜'+t['role'])
  note('题目',t['questions'][0]);note('参考答案',t['answer'])
  note('取证与必要加工',t['explain'])
  note('意义点','；'.join(t['points'])+'。')
  note('提示顺序','①'+t['hints'][0]+' ②'+t['hints'][1])
  if t['mode']=='written':note('教师核对',t['scoreGuide'])
  if t.get('oral'):
   for oq,oa in t['oral']:note('口头追问',oq+' '+oa)
  if id=='words-usage':
   for c in t['checks']:
    d.add_heading(c['question'],2)
    for i,o in enumerate(c['options']):p('ABCD'[i]+'．'+o+'　'+('✓ ' if i==c['correct'] else '')+c['why'][i])
  elif t.get('checks'):
   note('选项诊断',' '.join('题'+str(k+1)+'：'+ '；'.join('ABCD'[i]+' '+w for i,w in enumerate(c['why'])) for k,c in enumerate(t['checks'])))
  if id=='phrases':
   for n,b in enumerate(t['blanks'],1):p(str(n)+'．'+b['before']+' '+b['answer'])
  if id=='p5-link':note('常见误区','仅写“承上启下”而不写前后内容，只说明了作用名称；抄完整两句而不说明怎样连接，也未完成作用分析。')
  if id=='open-response':
   note('合理差异','可以选学校活动、社区生活、建筑或其他准确例子。不同立场和例子只要合理、具体、能解释多元与开放，均可接受。')
   note('不足答案','“牛车水体现多元开放，因为有很多文化。”只列地点并重复题意，需补充具体特色及接纳交流的表现。')
   note('订正指令','把自己的例子补充到可观察的具体内容；在例子后写一句“这体现……，因为……”。保留首次答案，另写订正。')
for node in d.element.xpath('.//w:pBdr'):
 node.getparent().remove(node)
for par in d.paragraphs:
 for run in par.runs:
  if '峇' in run.text:
   run.font.name='Huawentong WenKai';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Huawentong WenKai')
 par.paragraph_format.widow_control=True
 par.paragraph_format.keep_together=True
d.core_properties.title='多元而开放的新加坡：教学流程与参考答案';d.core_properties.author='华文通'
d.save(R/'downloads/teacher-guide.docx');print('DOCX written')
