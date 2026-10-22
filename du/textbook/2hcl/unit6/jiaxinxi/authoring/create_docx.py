from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]/'site/du/textbook/2hcl/unit6/jiaxinxi';L=json.loads((R/'lesson.json').read_text());D=Document();sec=D.sections[0];sec.page_height=Inches(11.69);sec.page_width=Inches(8.27);sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
K='Huawentong WenKai';S='Noto Serif CJK SC';styles=D.styles
for name in ['Normal','Body Text','Caption','List Bullet','List Number']:
 st=styles[name];st.font.name=S;st.font.size=Pt(10.5);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),S);st.paragraph_format.space_after=Pt(2);st.paragraph_format.line_spacing=1.08
for name,size in [('Title',25),('Heading 1',19),('Heading 2',13),('Heading 3',11)]:
 st=styles[name];st.font.name=S;st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('176962');st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),S)
head=sec.header.paragraphs[0];head.text='华文通  ·  中二高级华文  ·  单元六  ·  巩固课文';head.style='Caption'
foot=sec.footer.paragraphs[0];foot.alignment=2;foot.add_run('《别上假信息的当》  ·  ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
def para(text='',kai=False,style=None):
 p=D.add_paragraph(style=style);r=p.add_run(text)
 if kai:r.font.name=K;r.font.size=Pt(12);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),K)
 return p
def h(text,level=2):
 p=D.add_heading(text,level);p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(3)
def bullet(text):return para(text,style='List Bullet')
def page():D.add_page_break()
def table(headers,rows,widths=None):
 t=D.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for row in t.rows:
  for c in row.cells:
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(3)
 return t
D.add_heading('别上假信息的当',0);para('两课时教学流程与参考答案');para('中二高级华文 · 单元六 · 巩固课文 · 每课暂按50分钟');h('教学目标')
for x in L['teaching']['objectives']:bullet(x)
h('课堂流程');table(['课时','内容与活动','分钟'],L['teaching']['schedule'])
h('课堂组织');para('按段读文与看漫画 → 完成本段问题 → 核对证据与解释 → 进入下一段。第一段由教师示范，第二段师生共构，第三段开始独立解释。全文统整在第八段之后。');para('第一课书面输入：动荡原因、2012流言的作用。第二课书面输入：句式仿写、开放回应。其余使用点选、多句取证、排序或口头表达。时间不足时缩短汇报，不删课本原题与开放题。')
h('材料来源');para('本次上传：课本74—77页两张、教参三张、原漫画五页。保留全部40格；第31格为第5—6段小结，放在第6段之后。教参上网搜索和SLS视频仅作备用，不增加必做篇章或课后一周小测。')
page();h('使用与整合说明',1)
h('学生页面');para('同屏左侧当前段落，右侧相关任务。漫画逐格可直接阅读；点词可看释义并按需收藏。证据框可连续加入多句，也可粘贴与编辑。首次回答与订正分开保存。');para('客观题记录首次正确率；答错先给提示，等待5秒后可重试，不计速度分、不设排行榜。主观题只保存与显示参考答案，不用字数或关键词自动判分。开放题四部分必填。')
h('教师核对');para('先核对题干所问，再看答案是否保留意义点。相关原句只是证据；作用题还需解释证据与观点的关系。推断要有合理联系，不增添课文没有的事件细节。自检勾选不等于评分。')
h('旧内容去重');table(['旧资源','新版安排'],[['观点、刘邦、古今传播题','第1—3段整合；保留事例与观点解释。'],['三人成虎、危害排序','第5—6段整合；排序初始顺序打乱。'],['词语学习与阅读重复题','七道语境题随段，只做一次；词语本按需收藏。'],['方法题、停课情境','合为课本第三题，说明理由、应用与审慎处理。'],['旧随堂与精读入口','导向新学生／教师页；主目录只保留一个入口。']])
h('记录与下载');para('记录保存在同设备、同浏览器、同网站域名；不会自动上传教师或Google Sheet，没有账号云同步。可导出完整JSON、逐题CSV与文字记录。个人词语本跨已接入的课文保存，支持闪卡与导入导出。');para('旧本机记录与后台数据不删除；检查页可单独备份旧本机记录。原Netlify站不改写；不同域名的旧记录须在原站导出。离线与在线记录分开。整套下载附PPT、教师版、离线学习包、原图、字体与维护源文件。')
h('四项教学边界');bullet('刘邦杀蛇“可能确有其事”仍是不确定陈述，不能教成已证实史实。');bullet('第4段“以前”不能直接换成“古代”；图像也须查证。');bullet('多人重复转发不等于多个独立来源。');bullet('暂时查不到证实，不等于已经证明为假。先不转发，继续核实。')
classify={'p1-view':'直接信息：定位作者判断；区分现象与观点。','p2-example':'直接信息＋关系解释：区分事件与无依据的传言。','p2-cause':'合理推断：原文只写现象，需补出心理／需求与传播的联系。','p3-compare':'直接信息＋跨段整合：比较方式、数量、速度、范围。','p3-example':'直接信息＋关系解释：选择两处证据，联系本段观点。','p4-tech':'直接信息＋因果解释：技术修改与难辨真假。','p5-mislead':'直接信息＋概括道理：重复传播误导判断。','p6-harm':'直接信息＋因果整合：互信、合作、进步。','p7-method':'直接信息＋应用：提取方法并说明如何查证。','p7-sentence':'语言运用：目标与行动关系。','p8-ending':'结构关系：具体说明结尾如何回应开头。','whole-structure':'直接信息＋全文统整：采用课本结构名称。','open-response':'开放应用：方法合理、理由明确、例子具体、审慎处理。'}
for j,t in enumerate(L['tasks']):
 if j==0:page()
 h(t['title'],1);para('出处：'+t['source']+'｜'+t['role']);para(t['questions'][0],True)
 h('题型与加工');para(classify[t['id']]);h('参考答案');para(t['answer'],True);para(t['explain'])
 if t.get('evidence'):
  h('取证位置')
  for e in t['evidence']:para('第'+('、'.join(map(str,e['p'])))+'段；定位提示：'+'／'.join(e.get('requiredTerms',e['terms']))+'。须回到完整原句，不只抄关键词。')
 h('意义点与核对');para('；'.join(t['points'])+'。');para(t['scoreGuide'])
 if t.get('blanks'):
  h('逐空答案')
  for i,b in enumerate(t['blanks']):para(str(i+1)+'. '+b['before']+b['answer']+b['after'],True)
 if t.get('checks'):
  h('点选讲评')
  for c in t['checks']:
   para(c['question']+'　答案：'+'ABCD'[c['correct']],True)
   for i,o in enumerate(c['options']):para('ABCD'[i]+'. '+o+'｜'+c['why'][i])
 h('分层提示与误区');para('提示一：'+t['hints'][0]+'\n提示二：'+t['hints'][1]);para('不充分或错误：'+t['reject'])
 for a in t.get('alternatives',[]):para('可接受：'+a)
 for q,a in t.get('oral',[]):h('口头追问');para(q,True);para(a,True)
page();h('漫画对应与课本原题核对',1);table(['段落','原漫画','学习重点'],[[str(i+1),str(v[0])+'—'+str(v[-1]),L['paragraphIdeas'][i]]for i,v in enumerate(L['panels'])]);h('原题覆盖清单');table(['课本位置','新版任务'],[['第2段旁注','2b 动荡时为什么容易相信'],['第3段旁注','3a 古今传播有什么不同'],['第4段旁注','4 科技为什么使辨别更难'],['第5段旁注','5 重复为什么会误导人'],['第6段旁注','6 社会危害怎样形成'],['第7段旁注','7a 怎样减少假信息的危害'],['巩固练习一','7b 句式仿写'],['巩固练习二','9 课本结构表，九空'],['巩固练习三','10 开放回应，四项必填']]);h('维护资料');para('课文、题目、答案共用lesson.json生成；原图不改动。在线投影与下载PPT采用同一套页面。词语本采用平台共享格式，按年级、单元、课文归类。')
D.save(R/'downloads/teacher-guide.docx');print('docx saved')
