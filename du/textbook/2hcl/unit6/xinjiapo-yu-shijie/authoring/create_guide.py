import json,pathlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=pathlib.Path('/workspace/scratch/6e755f6b886c');O=R/'huawentong-home/du/textbook/2hcl/unit6/xinjiapo-yu-shijie';L=json.loads((O/'lesson.json').read_text());B=json.loads((O/'teaching-brief.json').read_text())
d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
 s=d.styles[name];s.font.name='Noto Serif CJK SC';s._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');s.font.color.rgb=RGBColor(0,0,0)
normal=d.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=Pt(16);normal.paragraph_format.space_after=Pt(6)
d.styles['Title'].paragraph_format.line_spacing=Pt(30);d.styles['Title'].font.size=Pt(23);d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 2'].font.size=Pt(13)
foot=sec.footer.paragraphs[0];foot.alignment=2;r=foot.add_run('华文通 · 中二高级华文 · ');r.font.size=Pt(9);f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
def p(s,kai=False):
 a=d.add_paragraph(s)
 if kai:
  for r in a.runs:r.font.name='LXGW WenKai TC';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'LXGW WenKai TC');r.font.size=Pt(12)
 return a
def heading(s):d.add_heading(s,1)
def page():d.add_page_break()
def table(headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for row in t.rows:
  trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(4);par.paragraph_format.line_spacing=Pt(14)
    for run in par.runs:run.font.size=Pt(10)
 return t

d.add_heading('新加坡与世界的联系',0);p('教师流程与答案｜中二高级华文 · 单元六 · 生活空间');p('1课时：按50分钟设计。课文、课本问题和教参据教师提供附件整理。');heading('教学重点');
for i,s in enumerate(B['goals']):p(f'{i+1}. {s}')
heading('课堂流程');table(['环节','分钟','行动与预期表现'],B['schedule'])
heading('材料整合');p('采用“精简鸡饭导入＋课本主线教学＋统一入口”。四页原有漫画保留不改，课文按8个自然段顺序教学。旧活动只留下切入情境，删去六国猜配对、缺少依据的产地判断与重复练习。课本四道理解题、四项词语释义和开放题全部保留。');p('学生只需为第2题及开放题写完整答案，其余采用点选、选词或口头表达。词语释义集中检查一次，个人词语本用于按需要收藏与复习。')
p('平台位置：中二 → 课文学习 → 单元六 → 生活空间 · 新加坡与世界的联系。')
page();heading('教学推进与差异化支持');table(['时间','教师行动','学生行动／观察点'],[
['0—3分钟','问：鸡饭怎样让你想到外地？再出示SFA与PUB两条资料。','口头提出猜想，区分整体资料与具体食材产地。'],
['3—17分钟','依次读第1—4段。用第1题示范“货物”与“用途”的区别。','找全丝绸、瓷器、香料；说出货物流通也带来文化交流。'],
['17—26分钟','读第5—6段，先释“枢纽”，让学生独立作第2题。','点句、加入多句证据，再把答案写完整；检查三个意义点。'],
['26—36分钟','读第7—8段。共同找“不但……也……”，再独立作第4题。','辨清共同身份与文化传承；分清科技手段、数码化方面与预期作用。'],
['36—45分钟','开放题：5分钟独写、2分钟分享、2分钟讲评。','一个课外方面＋具体例子＋解释联系＋回扣。'],
['45—50分钟','完成四项词语释义，检查订正及词语本。','记录需要复习的词；保存或导出本课学习记录。']])
heading('支持与提升');p('需要支持：提供句框“新加坡……，并且……”；先口头说再写，不替学生给出内容。开放题可用“我认为……方面。例如……。这说明……，因为……”组织。');p('进度较快：用不同措辞表达相同意思，说明为什么“用途”不能替代“货物”；在开放回应中把跨国双方与实际联系说得更清楚。不另加短文或重复小测。');p('若本校一课时为40分钟：导入2、古代10、现代15、开放9、词语与订正4分钟。减少重复朗读与展示人数，保留开放题独立写答。');heading('材料表述边界');p('教材的“600多个港口”及“2014年智慧国愿景”按原文保留，用于理解课文，不作为最新统计或全部完成的事实。漫画的参观情境、综合回顾画格不是正文新段落，也不充当答案证据。')
page();heading('第1题：古代贸易货物');t=L['tasks'][0];p(t['questions'][0],True);p('对应：第2—4段；第一段为总述。题型：直接信息，跨段整合。')
p('参考答案：'+t['answer'],True);p('证据：第二段写丝绸起源与运输；第三段写瓷器海上贸易；第四段写香料经新加坡传播。');p('思路：题干问“哪些货物”→每段找一种货物→并列列出三类。');p('四选项与解析：');
for i,(op,why) in enumerate(zip(t['checks'][0]['options'],t['checks'][0]['why'])):p('ABCD'[i]+'. '+op+'｜'+why)
p('核对：'+t['scoreGuide']);p('口头追问：除了货物，还有什么随着贸易传播？参考：第一段的东西方文化、第四段香料的使用方法。无需另写一题。')
heading('第2题：全球物流枢纽');t=L['tasks'][1];p(t['questions'][0],True);p('对应：第6段。题型：根据原文说明判断的依据；不是靠个人常识猜原因。');p('原文证据：'+L['paragraphs'][5],True);p('参考答案：'+t['answer'],True);p('意义点：主要港口之一；与全球600多个港口相连；提供一流服务。');p('核对：'+t['scoreGuide']);p('易错：'+t['reject']);p('允许保留原意的等义表达，如“连接世界各地许多港口”可体现联系广泛；如以数字作证，应准确，不把港口说成船只。')
page();heading('第3题：多元文化社会');t=L['tasks'][2];p(t['questions'][0],True);p('对应：第7段“不但……也……”句。题型：直接信息，找全两项做法。');p('原文证据：她不但注重加强国人的身份认同，也鼓励各种族传承自己的文化。',True);p('参考答案：'+t['answer'],True);p('核对：'+t['scoreGuide']);p('讲评：身份认同说明共同归属；文化传承保留各种族文化，两者并不互相排斥。“开放心胸接纳各地文化”是本段进一步说明的特质及作用，不替代题目要找的两项做法。');p('干扰项把“多元”误作文化完全相同、放弃共同身份或用外国文化取代本族文化；应回到原句修正。')
heading('第4题：智慧国的愿景');t=L['tasks'][3];p(t['questions'][0],True);p('对应：第8段第一句。题型：直接信息，提取并列方面。');p('参考答案：'+t['answer'],True);p('核对：'+t['scoreGuide']);p('“充分利用高科技”是手段；“社会、经济、公共服务数码化”是愿景方面；“国民受惠、加强联系”是预期作用。');
for i,(op,why) in enumerate(zip(t['checks'][0]['options'],t['checks'][0]['why'])):p('ABCD'[i]+'. '+op+'｜'+why)
p('教学中可让学生指着“推动”后面的三个词说答案，无需再抄写一遍。网页选择答案后才显示参考与解析。')
page();heading('开放题：提出课文以外的联系');t=L['tasks'][4];p(t['questions'][0],True);p(t['questions'][1],True);p('前一句为课本第三大题，后一追问为教师补充，目的在于让例子与看法之间的关系说清楚。开放题必做，但不要求唯一答案。')
p('参考示例：'+t['answer'],True);table(['核对点','1分的表现','不足的表现'],[['方面','提出课文以外的一个方面','只重复物流、文化与科技的小标题'],['例子','说清谁与哪里发生什么具体联系','只有“旅游很好”等笼统判断'],['解释','指出人员、活动或经验的跨国往来','只把例子再说一遍'],['回扣','解释能支持所提方面','例子与所提方面不相干']]);p('以上4分仅供教师课堂核对，各点独立判断；不以连接词数量或长度评分，不要求学生写亲身经历。旅游、跨国教育交流、国际体育活动等合理方向均可；具体机构、赛事与数字应核实。')
heading('反馈示例');p('学生：“我觉得旅游，因为新加坡有很多景点。”教师：“你提出了方面。谁从哪里来，或谁到哪里去？这种往来怎样体现新加坡与世界的联系？”');p('学生：“新加坡也参加体育比赛。”教师：“把跨国联系补清楚：参赛者与哪些地方的人在什么活动中接触？这为什么说明体育方面的联系？”');heading('网页处理');p('提交后保存首次答案，再展示参考供比较。网页只检查是否有填答，不判断观点和例子是否合理；自检勾选不等于得分。教师核对后让学生在订正框完善原答，保留首次作答。')
page();heading('课本词语检查与个人词语本');t=L['tasks'][5];table(['课本释义','对应词语','原文位置'],[[b['before'],b['answer'],loc] for b,loc in zip(t['blanks'],['第2段“奢侈品”','第6段小标题“全球物流枢纽”','第8段第一句','第8段第二句'])]);p('四项释义完整保留，只集中检查一次。“奢侈”与“奢侈品”分词条解释；“枢纽”明确标记来源为课文小标题，避免杜撰原句。');p('词语本沿用华文通统一记录，支持年级／单元／课文筛选、已掌握与待复习标记、例句、闪卡及导入导出。本课37个词提供释义，学生按实际需要收藏，不强迫全部收录。')
heading('错误后的教学支持');table(['任务','第一层：定位','第二层：筛选或解释'],[[t['title'],*t['hints']] for t in L['tasks']]);p('客观题先给提示，等待5秒再试；再次尝试后可看讲解，首次作答仍保留。不要将“整道客观任务首次全对比例”解读为逐分成绩。第2题与开放题不计入自动正确率。')
heading('记录与回读');p('学习记录可导出完整备份、逐题记录及文字记录；词语本单独导出。网站与离线文件各自保存于所用浏览器。换设备或清除浏览器资料前先备份。完整漫画供读后串联回顾，不增加新的必交任务。')
page();heading('图文对应与资料');table(['教学段','内容','原漫画格'],[[i+1,L['paragraphLabels'][i],'、'.join(map(str,L['panels'][i]))] for i in range(8)]);p('四页原图保持完整；网页仅按对应画格裁切展示。第2格为参观情境；第24、32格为综合回顾。它们保留在完整漫画中，不当作额外自然段。')
heading('来源与使用说明');p('1. 课文、阅读放大镜及教学参考：教师提供的四张教材截图。原课文8段按教材顺序转录。建议课堂量表由本配套编制，不标作官方答案分值。');
for i,s in enumerate(B['sources'][1:],2):p(f'{i}. {s["title"]}\n{s["url"]}\n用途：{s["purpose"]}。')
p('导入资料核对日期：2026年10月21日。食物统计用整体进口比例说明联系；水源资料仅说明部分原水进口，均不能推出某一盘饭或某一杯水的实际来源。');heading('配套文件');p('课程主页：统一入口与3分钟导入；互动学习：逐段图文和六个任务；教师专区：流程与答案；在线投影及可编辑PPT；教师Word；离线互动HTML；完整配套ZIP。原Claude活动链接仅在教师资料中留作旧版参考。')
for el in d.element.xpath('//w:pBdr'): el.getparent().remove(el)
for style in d.styles:
 for el in style.element.xpath('.//w:pBdr'): el.getparent().remove(el)
 for el in style.element.xpath('.//w:rFonts'):
  for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
   el.attrib.pop(qn('w:'+attr),None)
# Explicit font on every run to prevent theme fallback in Word.
for el in d.element.xpath('//w:r'):
 rp=el.find(qn('w:rPr'))
 if rp is None: rp=OxmlElement('w:rPr');el.insert(0,rp)
 rf=rp.find(qn('w:rFonts'))
 if rf is None: rf=OxmlElement('w:rFonts');rp.insert(0,rf)
 font=rf.get(qn('w:eastAsia'),'Noto Serif CJK SC')
 for attr in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+attr),font)
d.save(O/'downloads/teacher-guide.docx');print('saved',O/'downloads/teacher-guide.docx')
